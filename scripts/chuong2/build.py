# -*- coding: utf-8 -*-
"""Dựng Chương 2 bản chuẩn hóa dữ liệu, có biểu đồ Excel gốc nhúng trong Word.

Chạy từ thư mục gốc của kho:
    python3 scripts/chuong2/build.py

Đầu vào : Chuong_2_Thuc_trang_hoan_chinh mới.docx
Đầu ra  : Chuong_2_Thuc_trang_chuan_hoa_bieu_do.docx
          Du_lieu_bieu_do_Chuong_2.xlsx
"""
import copy
import io
import os
import re
import sys

import docx
import xlsxwriter
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bieu_do as B  # noqa: E402
import du_lieu as D  # noqa: E402
from excel_chart import ghi_bang, ve  # noqa: E402

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
VAO = os.path.join(GOC, "Chuong_2_Thuc_trang_hoan_chinh mới.docx")
RA_DOCX = os.path.join(GOC, "Chuong_2_Thuc_trang_chuan_hoa_bieu_do.docx")
RA_XLSX = os.path.join(GOC, "Du_lieu_bieu_do_Chuong_2.xlsx")

CT_CHART = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RT_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
EMU_CM = 360000
TEN_SHEET_NHUNG = "Du_lieu"


# ---------------------------------------------------------------------------
# Định dạng dùng chung cho workbook
# ---------------------------------------------------------------------------
def dinh_dang(wb):
    base = {"font_name": "Times New Roman", "font_size": 11, "valign": "vcenter"}
    f = {
        "tieu_de": wb.add_format({**base, "bold": True, "font_size": 13}),
        "nguon": wb.add_format({**base, "italic": True, "text_wrap": True, "font_size": 10}),
        "cot": wb.add_format({**base, "bold": True, "bg_color": "#DCE8F7", "border": 1,
                              "border_color": "#9FB7D6", "text_wrap": True, "align": "center"}),
        "o": wb.add_format({**base, "border": 1, "border_color": "#C9C8C3", "text_wrap": True}),
        "lien_ket": wb.add_format({**base, "font_color": "#1F5FB0", "underline": 1}),
        "thuong": wb.add_format({**base, "text_wrap": True}),
    }
    f["so"] = {k: wb.add_format({**base, "border": 1, "border_color": "#C9C8C3", "num_format": k})
               for k in ("0", "0.0", "0.00", "0.0%", "0%", "#,##0")}
    return f


def cao_cm(h):
    if h.get("cao"):
        return h["cao"]
    loai = h["bieu_do"]["loai"]
    if loai == "doughnut":
        return 7.0
    if loai == "bar":
        return 6.5 + 0.55 * len(h["dong"])
    return 9.0


# ---------------------------------------------------------------------------
# 1. Workbook tổng hợp: mỗi hình một sheet có bảng dữ liệu và biểu đồ
# ---------------------------------------------------------------------------
NHAT_KY = [
    ("Bảng 2.1, tổng các dòng", "Các dòng cộng lại được 246 người trong khi tổng là 252; dòng Các đơn vị khác gộp số "
     "liệu trình độ của Viện Nghiên cứu giáo dục và Chuyển giao tri thức nhưng không gộp số người",
     "Tách dòng Viện Nghiên cứu giáo dục và Chuyển giao tri thức 6 người; dòng Các đơn vị khác: 59 người, tiến sĩ 6, "
     "thạc sĩ 22, đại học 15", "Tai lieu thanh do/2026 DS.xlsx"),
    ("Mục 2.1.1, số lĩnh vực đào tạo", "Bảy lĩnh vực nhưng liệt kê chín", "Chín lĩnh vực",
     "Thông tư 74/2026/TT-BGDĐT, mã lĩnh vực 14, 22, 31, 34, 38, 48, 51, 72, 81"),
    ("Bảng 2.2, đề tài cấp cơ sở", "37 đề tài: 15, 0, 7, 7, 8", "38 đề tài: 6, 10, 7, 8, 7 theo năm trong mã số",
     "Tổng hợp đề tài KHCN cấp cơ sở của GV 2021-2025.pdf, 38 mã số"),
    ("Bảng 2.2, tham luận hội thảo quốc tế", "15 lượt, năm 2024 có 4", "16 lượt, năm 2024 có 5",
     "Thống kê báo cáo, tham luận Hội nghị, Hội thảo quốc tế 2021-2025.pdf, cột ngày tổ chức"),
    ("Bảng 2.2, đề tài cấp quốc gia", "Chưa tách theo năm", "2024: 1; 2025: 2",
     "Đề tài KHCN cấp quốc gia thực hiện giai đoạn 2021-2025.pdf, ngày quyết định phê duyệt kinh phí"),
    ("Tổng bản ghi và sản phẩm độc lập", "580 bản ghi; khoảng 470 sản phẩm độc lập",
     "582 bản ghi; khoảng 475 sản phẩm độc lập", "Hệ quả của hai dòng trên"),
    ("Kinh phí đề tài cấp cơ sở", "20 đề tài được cấp 374,25 triệu đồng; 15 quy đổi giờ; 2 tự tìm tài trợ",
     "19 đề tài được Trường cấp 424,75 triệu đồng; 16 quy đổi giờ; 3 tự tìm tài trợ",
     "Cột kinh phí thực hiện trong danh mục đề tài cấp cơ sở"),
    ("Xếp loại nghiệm thu", "19 Tốt, 16 Đạt, 1 Khá, 1 Xuất sắc, tổng 37", "1 Xuất sắc, 19 Tốt, 1 Khá, 17 Đạt, tổng 38",
     "Cột xếp loại trong danh mục đề tài cấp cơ sở"),
    ("Bảng 2.3 và chỉ số Herfindahl", "Viện Quản trị và Công nghệ 9 đề tài; Herfindahl 0,237",
     "10 đề tài, gồm đề tài do Khoa Sau đại học chủ trì; Herfindahl 0,238",
     "Quy tắc quy đổi đơn vị đã áp dụng cho danh mục giáo trình"),
    ("Bảng 2.4", "Thiếu quy định tại điểm b khoản 4 Điều 36 Quyết định 213; mức 50 triệu đồng được hiểu là mức thưởng",
     "Bổ sung quy định 30% tác giả, 20% đơn vị, 50% Quỹ nghiên cứu khoa học; 50 triệu đồng là mức kinh phí xét duyệt "
     "đề tài", "Quyết định 213/QĐ-ĐHTĐ, Điều 36; Quy chế chi tiêu nội bộ 2026, mục 6.3.5.1"),
    ("Mục 2.2.2, nhân sự Phòng Khoa học Công nghệ", "Chỉ nêu số lượng 2 người",
     "Bổ sung trình độ, chuyên ngành và việc trưởng phòng do Hiệu trưởng kiêm nhiệm", "Tai lieu thanh do/2026 DS.xlsx"),
    ("Bảng 2.6", "Phát sinh ở tám nhóm, xác lập ở bốn nhóm",
     "Phát sinh ở chín nhóm, có hoạt động xác lập ở bốn nhóm, có văn bằng ở ba nhóm", "Đếm lại từ chính Bảng 2.6"),
    ("Bảng 2.7, dòng tổng và đoạn bình luận", "5 thương hiệu, 5 hợp tác, 2 nghiên cứu",
     "4 thương hiệu, 6 hợp tác, 2 nghiên cứu", "Đếm lại từ chính Bảng 2.7"),
    ("Bảng 2.8, năm của ba đề tài đầu", "2021", "2022", "Mã số 37-2022, 39-2022, 40-2022"),
    ("Liên hệ đề tài và tài sản trí tuệ", "Không có điểm giao nhau ngoài trường hợp lá Quế hoa",
     "Bổ sung liên hệ chưa được xác nhận giữa đề tài tinh dầu Citrus grandis năm 2023 và kiểu dáng công nghiệp "
     "tinh dầu vỏ bưởi đào năm 2024", "IP_Master_Dataset v08, sheet 13B_PROTECTION_FULL và 16_VERIFICATION_FINDINGS"),
    ("Mục 2.4.1, giáo trình gắn học phần", "61 trên 87 tài liệu, tương đương 70,1%, ghi số tín chỉ",
     "Toàn bộ 61 tài liệu giai đoạn 2022 - 2025 ghi số tín chỉ; danh mục năm 2021 không có trường số tín chỉ",
     "Thống kê giáo trình, tài liệu giảng dạy biên soạn 2021-2025.pdf"),
    ("Mục 2.1.5, danh mục hội thảo cấp trường", "Không trích xuất được nội dung",
     "Trích xuất được 279 lượt: 49, 50, 45, 60, 74 và 1 bản ghi sai năm tổ chức",
     "Thống kê báo cáo, tham luận Hội thảo cấp trường 2021-2025.pdf"),
    ("Mục 2.5.3", "Ghi sáu nguyên nhân nhưng liệt kê bảy; chưa tách khách quan và chủ quan",
     "Tách hai nhóm nguyên nhân; nhóm chủ quan gồm bảy nguyên nhân", "Quy tắc tại GEMINI.md"),
    ("Quy chế quản trị tài sản trí tuệ năm 2024", "Chưa được khai thác trong Chương 2",
     "Bổ sung vào Mục 2.2.1, 2.2.2, 2.3.1, 2.5.1, 2.5.3, Bảng 2.4 và Bảng 2.6",
     "Tai lieu thanh do/QC SHTT ĐHTĐ.docx; bản được cung cấp chưa ghi số và ngày ban hành"),
    ("Điều 34 Quyết định 213 và giải pháp hữu ích", "Khẳng định Điều 34 không liệt kê giải pháp hữu ích và dùng làm "
     "một nguyên nhân", "Điều 34 có liệt kê bằng độc quyền giải pháp hữu ích; sửa Bảng 2.6, Mục 2.3.1, 2.5.3 và lời bình "
     "Hình 2.14", "02. QĐ 213_QĐ-ĐHTĐ, Điều 34 khoản 2"),
    ("Số hiệu văn bản 213", "Quyết định số 213/QĐ-ĐHTĐ năm 2021",
     "Quy chế ban hành ngày 28/12/2021 kèm văn bản số 213/QĐ-ĐHTĐ; bản quy chế ghi là Nghị quyết, Kế hoạch 07 ghi Nghị "
     "quyết số 213/NQ-HĐT-ĐHTĐ. Đề nghị thống nhất cách gọi", "Trang bìa Quy chế 213; Kế hoạch 07/KH-ĐHTĐ"),
    ("Mục 2.1.5, độ phủ khai báo", "Hồ sơ khai 28 bài báo; 25 trên 28 tìm thấy; độ phủ 26,2%",
     "26 đề tài khai 32 công bố; độ phủ khoảng 29,9%; bỏ tỷ lệ 89,3% vì không tái lập được từ tài liệu gốc",
     "Cột sản phẩm nghiệm thu, danh mục đề tài cấp cơ sở"),
    ("Mục 2.5.2, ý thứ sáu", "Ba đề tài ghi rõ không có bài báo như cam kết",
     "Danh mục gốc không có ghi chú này; thay bằng: 7 đề tài chỉ có báo cáo tổng kết, 5 trong số đó xếp loại Tốt",
     "Danh mục đề tài cấp cơ sở"),
    ("Mục 2.2.3, lệ phí đăng ký", "Chỉ nêu không có dòng chi",
     "Bổ sung: Điều 35 giao Phòng KHCN nộp lệ phí nhưng Điều 38 không có mục chi tương ứng", "Quyết định 213, Điều 35, 38"),
    ("Kế hoạch 07/KH-ĐHTĐ ngày 01/7/2024", "Chưa sử dụng; tệp là bản quét",
     "Nhận dạng chữ và đối chiếu ảnh gốc: số liệu 2021 - 2023 khớp; bổ sung Hình 2.16 so sánh kế hoạch và thực hiện",
     "Tai lieu thanh do/KH hoạt động KHCN giai đoạn 2024-2028.pdf"),
    ("Tham luận hội thảo quốc tế 2022 - 2023", "",
     "Danh mục theo ngày tổ chức: 2 và 3; Kế hoạch 07 ghi 1 và 4. Chương 2 dùng số liệu danh mục",
     "Danh mục tham luận quốc tế; Kế hoạch 07, Bảng 4"),
    ("Kiểm chứng phép so khớp tác giả", "Chưa kiểm chứng",
     "Tái lập độc lập: 256 trên 405 bài khớp; Viện Y - Dược 55, Viện QT và CN 57, Viện NCGD và CGTT 17 bài; 89 người có "
     "bài, Gini 0,828; chênh 1 - 3 đơn vị do 5 họ tên trùng nhau. Giữ số liệu Chương 2",
     "2026 DS.xlsx và hai danh mục bài báo"),
    ("Nội dung chưa có tài liệu gốc trong thư mục", "",
     "Điều lệ Quỹ Ngô Xuân Độ; hợp đồng của Viện NCGD và CGTT; tư cách thành viên Mạng lưới Trung tâm Hỗ trợ công nghệ "
     "và đổi mới sáng tạo từ năm 2023; đăng ký hoạt động khoa học công nghệ của Viện; sơ đồ tổ chức ngày 16/6/2026. Giữ "
     "nguyên, cần bổ sung minh chứng", ""),
    ("Tệp Du_lieu_thong_ke_KHCN_Thanh_Do_2021_2025.xlsx", "Bài báo trong nước theo năm: 54, 50, 33, 44, 109",
     "Danh mục gốc: 17, 42, 50, 72, 109. Không dùng tệp này để trích dẫn",
     "Thống kê bài báo đăng tạp chí khoa học trong nước 2021-2025.pdf"),
]


def dung_workbook():
    wb = xlsxwriter.Workbook(RA_XLSX)
    f = dinh_dang(wb)

    ml = wb.add_worksheet("Muc_luc")
    ml.set_column(0, 0, 10)
    ml.set_column(1, 1, 95)
    ml.set_column(2, 2, 24)
    ml.write(0, 0, "DỮ LIỆU VÀ BIỂU ĐỒ CHƯƠNG 2 - THỰC TRẠNG QUẢN LÝ QUYỀN SỞ HỮU TRÍ TUỆ TẠI TRƯỜNG ĐẠI HỌC THÀNH ĐÔ",
             f["tieu_de"])
    ml.write(1, 0, "Mỗi hình có một sheet riêng gồm bảng dữ liệu và biểu đồ Excel. Sửa số liệu trong bảng thì biểu đồ tự "
                   "cập nhật. Biểu đồ trong tệp Word được nhúng kèm dữ liệu, sửa bằng lệnh Chỉnh sửa dữ liệu của Word.",
             f["thuong"])
    ml.write_row(3, 0, ["Hình", "Tên hình", "Sheet"], f["cot"])
    for i, h in enumerate(B.HINH):
        ml.write(4 + i, 0, f"Hình 2.{h['so']}", f["o"])
        ml.write(4 + i, 1, h["tieu_de"], f["o"])
        ml.write_url(4 + i, 2, f"internal:'{h['id']}'!A1", f["lien_ket"], h["id"])
    r = 5 + len(B.HINH)
    for ten, mo_ta in [("Nhat_ky_chuan_hoa", "Các chỉnh sửa số liệu so với bản Chương 2 trước"),
                       ("DL_De_tai", "Danh mục 38 đề tài cấp cơ sở đã chuẩn hóa, không gồm thông tin cá nhân"),
                       ("DL_TSTT", "Danh mục 12 tài sản trí tuệ thuộc sở hữu của Nhà trường"),
                       ("DL_Tieu_chi", "Đánh giá khả năng tính toán 16 tiêu chí tại Mục 1.4.2")]:
        ml.write(r, 1, mo_ta, f["o"])
        ml.write_url(r, 2, f"internal:'{ten}'!A1", f["lien_ket"], ten)
        r += 1

    for h in B.HINH:
        ws = wb.add_worksheet(h["id"])
        ws.set_landscape()
        ws.set_paper(9)
        ws.fit_to_pages(1, 0)
        ws.set_column(0, 0, 42)
        ws.set_column(1, len(h["cot"]) - 1, 17)
        ws.set_row(3, 45)
        ws.write(0, 0, f"Hình 2.{h['so']}. {h['tieu_de']}", f["tieu_de"])
        ws.merge_range(1, 0, 1, max(5, len(h["cot"]) - 1), h["nguon"], f["nguon"])
        ws.set_row(1, 42)
        r1, r2 = ghi_bang(ws, 3, 0, h, f["cot"], f["o"], f["so"])
        ch = ve(wb, h["id"], h, 3, 0, r1, r2, tieu_de=f"Hình 2.{h['so']}. {h['tieu_de']}")
        ch.set_size({"width": 900, "height": int(cao_cm(h) / 9 * 520) + 40})
        ws.insert_chart(r2 + 2, 0, ch)

    nk = wb.add_worksheet("Nhat_ky_chuan_hoa")
    nk.set_column(0, 0, 32)
    nk.set_column(1, 2, 55)
    nk.set_column(3, 3, 50)
    nk.write(0, 0, "NHẬT KÝ CHUẨN HÓA SỐ LIỆU CHƯƠNG 2", f["tieu_de"])
    nk.write_row(2, 0, ["Nội dung", "Bản trước", "Giá trị chuẩn hóa", "Căn cứ đối chiếu"], f["cot"])
    for i, row in enumerate(NHAT_KY):
        nk.write_row(3 + i, 0, row, f["o"])

    dt = wb.add_worksheet("DL_De_tai")
    cot = ["Mã số", "Năm", "Đơn vị quy đổi", "Xếp loại", "Hình thức kinh phí", "Kinh phí Trường cấp (triệu đồng)",
           "Sản phẩm đủ điều kiện xác lập quyền", "Nhóm quyền có thể xác lập", "Tình trạng nộp đơn"]
    dt.set_column(0, 1, 10)
    dt.set_column(2, 4, 20)
    dt.set_column(5, 5, 16)
    dt.set_column(6, 8, 40)
    dt.write(0, 0, "DANH MỤC ĐỀ TÀI CẤP CƠ SỞ 2021 - 2025 ĐÃ CHUẨN HÓA", f["tieu_de"])
    dt.write_row(2, 0, cot, f["cot"])
    for i, d in enumerate(D.DE_TAI):
        dt.write_row(3 + i, 0, [d[0], d[1], d[2], d[3], d[4]], f["o"])
        dt.write_number(3 + i, 5, d[5], f["so"]["0.00"])
        dt.write_row(3 + i, 6, [d[6], d[7], d[8]], f["o"])
    n = 3 + len(D.DE_TAI)
    dt.write(n, 4, "Tổng", f["cot"])
    dt.write_formula(n, 5, f"=SUM(F4:F{n})", f["so"]["0.00"], B.kp_tong)
    dt.write(n + 2, 0, "Ghi chú: đề tài 14-2024 tự tìm nguồn tài trợ 19,5 triệu đồng, không tính vào kinh phí Trường "
                       "cấp; đề tài 09-2025 đề nghị cấp thêm 20 triệu đồng, chưa tính.", f["thuong"])

    ts = wb.add_worksheet("DL_TSTT")
    ts.set_column(0, 0, 22)
    ts.set_column(1, 1, 60)
    ts.set_column(2, 2, 8)
    ts.set_column(3, 5, 28)
    ts.write(0, 0, "TÀI SẢN TRÍ TUỆ THUỘC SỞ HỮU CỦA TRƯỜNG ĐẠI HỌC THÀNH ĐÔ", f["tieu_de"])
    ts.write_row(2, 0, ["Loại hình", "Tên tài sản", "Năm", "Trạng thái", "Cơ cấu sở hữu", "Nguồn hình thành"], f["cot"])
    for i, t in enumerate(D.TSTT):
        ts.write_row(3 + i, 0, t, f["o"])

    tc = wb.add_worksheet("DL_Tieu_chi")
    tc.set_column(0, 0, 14)
    tc.set_column(1, 1, 60)
    tc.set_column(2, 2, 22)
    tc.set_column(3, 3, 70)
    tc.write(0, 0, "KHẢ NĂNG TÍNH TOÁN BỘ TIÊU CHÍ TẠI MỤC 1.4.2", f["tieu_de"])
    tc.write_row(2, 0, ["Nhóm", "Tiêu chí", "Mức độ", "Căn cứ"], f["cot"])
    ten_muc = {"T": "Tính được đầy đủ", "M": "Tính được một phần", "C": "Chưa tính được"}
    for i, t in enumerate(D.TIEU_CHI):
        tc.write_row(3 + i, 0, [t[0], t[1], ten_muc[t[2]], t[3]], f["o"])
    wb.close()


# ---------------------------------------------------------------------------
# 2. Biểu đồ nhúng: tạo XML biểu đồ và workbook dữ liệu cho từng hình
# ---------------------------------------------------------------------------
def xlsx_bytes(h, co_bieu_do):
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    f = dinh_dang(wb)
    ws = wb.add_worksheet(TEN_SHEET_NHUNG)
    ws.set_column(0, 0, 40)
    ws.set_column(1, len(h["cot"]) - 1, 16)
    r1, r2 = ghi_bang(ws, 0, 0, h, f["cot"], f["o"], f["so"])
    ws.write(r2 + 2, 0, f"Hình 2.{h['so']}. {h['tieu_de']}", f["tieu_de"])
    ws.write(r2 + 3, 0, h["nguon"], f["nguon"])
    if co_bieu_do:
        ch = ve(wb, TEN_SHEET_NHUNG, h, 0, 0, r1, r2)
        ws.insert_chart(r2 + 5, 0, ch)
    wb.close()
    return buf.getvalue()


C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
THU_TU_NHAN = ["showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize"]


def chuan_hoa_nhan(xml):
    """Ghi tường minh mọi cờ hiển thị của nhãn dữ liệu để Word, Excel và LibreOffice hiển thị như nhau."""
    root = etree.fromstring(xml.encode("utf-8"))
    for dl in root.iter(f"{{{C_NS}}}dLbls", f"{{{C_NS}}}dLbl"):
        if dl.find(f"{{{C_NS}}}delete") is not None:
            continue
        gia_tri = {k: "0" for k in THU_TU_NHAN}
        for k in THU_TU_NHAN:
            e = dl.find(f"{{{C_NS}}}{k}")
            if e is not None:
                gia_tri[k] = e.get("val", "1")
                dl.remove(e)
        # chèn sau numFmt, spPr, txPr, dLblPos (nếu có), trước separator và showLeaderLines
        vi_tri = 0
        for i, e in enumerate(dl):
            if etree.QName(e).localname in ("idx", "layout", "tx", "numFmt", "spPr", "txPr", "dLblPos"):
                vi_tri = i + 1
        for j, k in enumerate(THU_TU_NHAN):
            e = etree.Element(f"{{{C_NS}}}{k}")
            e.set("val", gia_tri[k])
            dl.insert(vi_tri + j, e)
    # Đường gấp khúc: ghi tường minh smooth = 0, vì một số phần mềm mặc định làm mượt khi thiếu thẻ này
    for lc in root.iter(f"{{{C_NS}}}lineChart"):
        for ser in lc.findall(f"{{{C_NS}}}ser"):
            if ser.find(f"{{{C_NS}}}smooth") is None:
                e = etree.Element(f"{{{C_NS}}}smooth")
                e.set("val", "0")
                ext = ser.find(f"{{{C_NS}}}extLst")
                if ext is not None:
                    ext.addprevious(e)
                else:
                    ser.append(e)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True).decode("utf-8")


def chart_xml(h, rid_goi):
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(xlsx_bytes(h, True)))
    xml = chuan_hoa_nhan(z.read("xl/charts/chart1.xml").decode("utf-8"))
    xml = xml.replace('<c:lang val="en-US"/>', '<c:lang val="vi-VN"/><c:roundedCorners val="0"/>', 1)
    ext = f'<c:externalData r:id="{rid_goi}"><c:autoUpdate val="0"/></c:externalData>'
    if "<c:printSettings" in xml:
        xml = xml.replace("<c:printSettings", ext + "<c:printSettings", 1)
    else:
        xml = xml.replace("</c:chartSpace>", ext + "</c:chartSpace>", 1)
    return xml.encode("utf-8")


INLINE = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:rPr><w:noProof/></w:rPr><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
    '<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:docPr id="{id}" name="Biểu đồ {id}" descr="{descr}"/><wp:cNvGraphicFramePr/>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart r:id="{rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
)


def nhung_bieu_do(doc, h, n):
    pkg = doc.part.package
    goi = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx"), CT_XLSX, xlsx_bytes(h, False), pkg)
    bd = Part(PackURI(f"/word/charts/chart{n}.xml"), CT_CHART, b"", pkg)
    rid_goi = bd.relate_to(goi, RT_PACKAGE)
    bd._blob = chart_xml(h, rid_goi)
    rid = doc.part.relate_to(bd, RT.CHART)
    xml = INLINE.format(cx=int(15.5 * EMU_CM), cy=int(cao_cm(h) * EMU_CM), id=500 + n, rid=rid,
                        descr=f"Hình 2.{h['so']}. {h['tieu_de']}")
    return etree.fromstring(xml)


# ---------------------------------------------------------------------------
# 3. Chỉnh sửa tệp Word
# ---------------------------------------------------------------------------
def doan_than(doc):
    return [p for p in doc.paragraphs]


def dat_chu(p, text, bold=None, italic=None):
    runs = p.runs
    if not runs:
        r = p.add_run(text)
    else:
        r = runs[0]
        r.text = text
        for x in runs[1:]:
            x._r.getparent().remove(x._r)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    return p


def tim(doc, dau):
    ds = [p for p in doc.paragraphs if p.text.strip().startswith(dau)]
    assert len(ds) == 1, f"Tìm thấy {len(ds)} đoạn bắt đầu bằng: {dau}"
    return ds[0]


def thay(doc, cu, moi, toan_bo=False):
    ds = [p for p in doc.paragraphs if cu in p.text]
    if toan_bo:
        assert ds, f"Không tìm thấy: {cu}"
    else:
        assert len(ds) == 1, f"Tìm thấy {len(ds)} đoạn chứa: {cu}"
    for p in ds:
        dat_chu(p, p.text.replace(cu, moi))


def o(bang, hang, cot, text):
    c = bang.rows[hang].cells[cot]
    dat_chu(c.paragraphs[0], text)
    for p in c.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def them_hang(bang, sau_hang, gia_tri):
    tr = copy.deepcopy(bang.rows[sau_hang]._tr)
    bang.rows[sau_hang]._tr.addnext(tr)
    moi = bang.rows[sau_hang + 1]
    for j, v in enumerate(gia_tri):
        dat_chu(moi.cells[j].paragraphs[0], v)


def doan_moi(sau_el, mau, text, bold=None, italic=None, giu_voi_sau=False):
    el = copy.deepcopy(mau._p)
    for r in el.findall(qn("w:r")):
        el.remove(r)
    for tag in ("w:hyperlink", "w:bookmarkStart", "w:bookmarkEnd"):
        for x in el.findall(qn(tag)):
            el.remove(x)
    mau_r = mau.runs[0]._r if mau.runs else None
    r = copy.deepcopy(mau_r) if mau_r is not None else etree.SubElement(el, qn("w:r"))
    for t in r.findall(qn("w:t")):
        r.remove(t)
    el.append(r)
    sau_el.addnext(el)
    p = docx.text.paragraph.Paragraph(el, mau._parent)
    run = p.runs[0]
    run.text = text
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if giu_voi_sau:
        p.paragraph_format.keep_with_next = True
    return p


def dat_cap_de_muc(doc):
    for p in doc.paragraphs:
        t = p.text.strip()
        lvl = None
        if t == "CHƯƠNG 2":
            lvl = 0
        elif re.match(r"^2\.\d\. ", t) and p.runs and p.runs[0].bold:
            lvl = 1
        elif re.match(r"^2\.\d\.\d\. ", t) and p.runs and p.runs[0].bold:
            lvl = 2
        if lvl is None:
            continue
        ppr = p._p.get_or_add_pPr()
        for x in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(x)
        ol = etree.SubElement(ppr, qn("w:outlineLvl"))
        ol.set(qn("w:val"), str(lvl))
        # outlineLvl phải đứng trước rPr trong pPr
        rpr = ppr.find(qn("w:rPr"))
        if rpr is not None:
            ppr.remove(ol)
            rpr.addprevious(ol)
        p.paragraph_format.keep_with_next = True


def sua_so_lieu(doc):
    T = doc.tables
    # --- 2.1.1
    thay(doc, "trải trên bảy lĩnh vực đào tạo cấp hai", "trải trên chín lĩnh vực đào tạo")
    p = tim(doc, "Tính đến năm 2026, tổng nhân sự của Nhà trường là 252 người")
    dat_chu(p, "Tính đến năm 2026, tổng nhân sự của Nhà trường là 252 người. Nhóm có trình độ tiến sĩ và tương đương "
               "gồm 94 người, trong đó có 2 giáo sư và 23 phó giáo sư; nhóm có trình độ thạc sĩ và tương đương gồm 94 "
               "người. Cơ cấu nhân lực theo đơn vị được trình bày tại Bảng 2.1.")
    o(T[0], 7, 0, "Các đơn vị khác")
    for j, v in enumerate(["59", "0", "2", "6", "22", "15", "16"], 1):
        o(T[0], 7, j, v)
    them_hang(T[0], 6, ["Viện Nghiên cứu giáo dục và Chuyển giao tri thức", "6", "0", "0", "1", "1", "4", "0"])
    thay(doc, "Đây là mẫu số được sử dụng cho các chỉ số năng suất nghiên cứu trong toàn bộ chương này.",
         "Đây là mẫu số của các chỉ số năng suất nghiên cứu ở cấp toàn trường. Ở cấp đơn vị, Mục 2.1.4 sử dụng tổng "
         "nhân sự làm mẫu số vì phép so khớp tác giả được thực hiện trên toàn bộ danh sách nhân sự.")
    # --- 2.1.2
    thay(doc, "ghi nhận 580 bản ghi sản phẩm khoa học", f"ghi nhận {B.tong_ban_ghi} bản ghi sản phẩm khoa học")
    b = T[1]
    for i, (ten, v) in enumerate(D.SAN_PHAM):
        hang = {0: 1, 1: 2, 2: 3, 3: 4, 4: 5, 5: 6, 6: 8}[i]
        assert b.rows[hang].cells[0].text.strip().lower().startswith(ten.lower()[:12]), (ten, b.rows[hang].cells[0].text)
        for j, x in enumerate(v):
            o(b, hang, 1 + j, str(x) if x else "-")
        o(b, hang, 6, str(sum(v)))
    for j in range(5):
        o(b, 9, 1 + j, str(B.tong_nam[j]))
    o(b, 9, 6, str(B.tong_ban_ghi))
    thay(doc, "Danh mục không ghi nhận đề tài cấp cơ sở nào trong năm 2022.",
         "Đề tài cấp cơ sở được xếp theo năm ghi trong mã số đề tài; đề tài cấp quốc gia được xếp theo năm phê duyệt "
         "kinh phí; tham luận hội thảo quốc tế được xếp theo ngày tổ chức. Số liệu bài báo, giáo trình và đề tài cấp cơ sở "
         "giai đoạn 2021 - 2023 khớp với số liệu tổng kết tại Kế hoạch số 07/KH-ĐHTĐ ngày 01 tháng 7 năm 2024 của Nhà "
         "trường.")
    thay(doc, "Con số 580 chỉ phản ánh", f"Con số {B.tong_ban_ghi} chỉ phản ánh")
    thay(doc, "quy mô thực tế là khoảng 470 sản phẩm độc lập", f"quy mô thực tế là khoảng {B.doc_lap} sản phẩm độc lập")
    # --- 2.1.4
    o(T[2], 2, 2, "10")
    o(T[2], 2, 5, "21,3")
    thay(doc, "cột này chỉ dùng để so sánh tương đối giữa các đơn vị.",
         "cột này chỉ dùng để so sánh tương đối giữa các đơn vị. Đề tài do Khoa Sau đại học chủ trì được tính cho Viện "
         "Quản trị và Công nghệ; một đề tài do Phòng Khoa học Công nghệ chủ trì không được gán vào đơn vị nào.")
    thay(doc, "là 0,237 đối với đề tài", "là 0,238 đối với đề tài")
    # --- ngôi xưng
    thay(doc, "Tổng hợp của tác giả", "Tổng hợp của nhóm nghiên cứu", toan_bo=True)
    thay(doc, "tác giả tiến hành", "nhóm nghiên cứu tiến hành", toan_bo=True)
    # --- 2.1.5
    thay(doc, "Danh mục tham luận hội thảo cấp trường được lưu dưới dạng tệp không trích xuất được nội dung.",
         "Danh mục tham luận hội thảo cấp trường gồm 279 lượt báo cáo nhưng có bản ghi sai năm tổ chức, cụ thể một báo "
         "cáo được ghi ngày tổ chức vào năm 1905. Danh mục giáo trình năm 2021 dùng biểu mẫu khác với các năm sau và "
         "không có trường số tín chỉ.")
    thay(doc, "theo dõi tài sản trí tuệ đã được đăng ký hoặc đã được cấp văn bằng bảo hộ.",
         "theo dõi tài sản trí tuệ đã được đăng ký hoặc đã được cấp văn bằng bảo hộ, dù Điều 11 Quy chế quản trị tài sản "
         "trí tuệ năm 2024 giao Phòng nhiệm vụ lập hồ sơ thống kê, theo dõi tài sản trí tuệ thuộc sở hữu của Nhà trường.")
    p = tim(doc, "Thứ hai, mức độ khai báo sản phẩm")
    dat_chu(p, "Thứ hai, mức độ khai báo sản phẩm trong hồ sơ nghiệm thu đề tài thấp. Hồ sơ nghiệm thu của 26 trên 38 đề "
               "tài khai tổng cộng 32 công bố, trong khi phép đối chiếu tên chủ nhiệm đề tài với danh sách tác giả gắn được "
               "107 bài báo với các đề tài, tức độ phủ khai báo chỉ khoảng 29,9%. Phần lớn công bố có liên quan đến đề tài "
               "vì vậy không được phản ánh trong hồ sơ nghiệm thu. Nếu bài báo, là loại sản phẩm dễ nhận diện nhất, còn "
               "không được khai đủ, thì các sản phẩm dạng công thức, quy trình hoặc sản phẩm vật chất càng ít khả năng được "
               "ghi nhận đầy đủ.")
    # --- 2.2.1
    thay(doc, "Văn bản nền tảng điều chỉnh hoạt động sở hữu trí tuệ trong Nhà trường là Quy chế hoạt động khoa học công "
              "nghệ ban hành kèm Quyết định số 213/QĐ-ĐHTĐ năm 2021.",
         "Văn bản đầu tiên điều chỉnh hoạt động sở hữu trí tuệ trong Nhà trường là Quy chế hoạt động khoa học công nghệ "
         "ban hành ngày 28 tháng 12 năm 2021 kèm theo văn bản số 213/QĐ-ĐHTĐ, sau đây gọi là Quyết định 213.")
    thay(doc, "liệt kê tương đối đầy đủ từ tên trường, logo, nhãn hiệu đến công trình khoa học, chương trình đào tạo, ngân "
              "hàng đề thi, phần mềm, mô hình thực hành, sáng chế và quy trình công nghệ.",
         "liệt kê tương đối đầy đủ từ tên trường, logo, nhãn hiệu, bằng độc quyền sáng chế, giải pháp hữu ích, kiểu dáng "
         "công nghiệp đến công trình khoa học, chương trình đào tạo, ngân hàng đề thi, bản ghi âm, ghi hình, phần mềm, mô "
         "hình thực hành và quy trình công nghệ.")
    thay(doc, "Bên cạnh đó, Nhà trường còn có hai văn bản khác", "Bên cạnh hai quy chế trên, Nhà trường còn có hai văn bản khác")
    thay(doc, "Điều 36 quy định công thức phân chia lợi ích từ thương mại hóa.",
         "Điều 36 quy định hai công thức phân chia lợi ích từ thương mại hóa, tùy theo nguồn hình thành tài sản.")
    thay(doc, "quy định mức thưởng cho đề tài có đăng ký sở hữu trí tuệ, giải pháp hữu ích hoặc độc quyền sáng chế, kèm "
              "định mức quy đổi giờ nghiên cứu khoa học cho từng loại văn bằng.",
         "quy định mức kinh phí xét duyệt cho loại đề tài có đăng ký sở hữu trí tuệ, kèm điều kiện trích kinh phí chuyển "
         "giao công nghệ về Trường và định mức quy đổi giờ nghiên cứu khoa học cho từng loại văn bằng.")
    thay(doc, "Ba văn bản này quy định ba công thức chia lợi ích khác nhau cho cùng một loại tài sản, như trình bày tại "
              "Bảng 2.4.",
         "Bốn văn bản này chứa năm quy định khác nhau về phân chia lợi ích từ tài sản trí tuệ, như trình bày tại Bảng "
         "2.4.")
    thay(doc, "Bảng 2.4. Ba công thức chia lợi ích từ tài sản trí tuệ trong nội bộ Trường Đại học Thành Đô",
         "Bảng 2.4. Các quy định về phân chia lợi ích từ tài sản trí tuệ trong nội bộ Trường Đại học Thành Đô")
    b = T[3]
    o(b, 1, 0, "Quyết định số 213/QĐ-ĐHTĐ, điểm a khoản 4 Điều 36")
    them_hang(b, 1, ["Quyết định số 213/QĐ-ĐHTĐ, điểm b khoản 4 Điều 36", "2021",
                     "Tài sản trí tuệ thuộc sở hữu của Trường được chuyển giao",
                     "30% tác giả, 20% khoa hoặc đơn vị có tác giả, 50% Quỹ nghiên cứu khoa học của Trường",
                     "Không đặt trần"])
    them_hang(b, 2, ["Quy chế quản trị tài sản trí tuệ, Điều 13", "2024",
                     "Tài sản trí tuệ của Trường khi các bên không có thỏa thuận",
                     "Sau khi trừ chi phí, Hiệu trưởng quyết định tỷ lệ theo tham mưu của Phòng Khoa học Công nghệ, Bộ "
                     "phận Pháp chế và Phòng Tài chính - Kế toán",
                     "Không quy định"])
    o(b, 4, 3, "Kinh phí xét duyệt cho đề tài, với điều kiện có chứng nhận đăng ký thành công; trích 50% kinh phí "
               "chuyển giao công nghệ về Trường, không nêu phần của tác giả")
    o(b, 4, 4, "50 triệu đồng kinh phí đề tài")
    thay(doc, "Tổng hợp của nhóm nghiên cứu từ ba văn bản nội bộ",
         "Tổng hợp của nhóm nghiên cứu từ bốn văn bản nội bộ")
    thay(doc, "Ba công thức này chưa được rà soát để bảo đảm tính thống nhất. Một giảng viên có sản phẩm đủ điều kiện "
              "bảo hộ không thể xác định phần lợi ích của mình theo văn bản nào: 30% theo Quyết định 213, một khoản "
              "thưởng tối đa 50 triệu đồng theo Quy chế chi tiêu nội bộ, hay 50% giảm dần xuống 20% theo Điều lệ Quỹ.",
         "Năm quy định này chưa được rà soát để bảo đảm tính thống nhất. Một giảng viên có sản phẩm đủ điều kiện bảo hộ "
         "không thể xác định phần lợi ích của mình theo quy định nào: 30% có trần 100 triệu đồng theo điểm a hay 30% "
         "không đặt trần theo điểm b của cùng khoản 4 Điều 36 Quyết định 213, một tỷ lệ do Hiệu trưởng quyết định theo "
         "Điều 13 Quy chế quản trị tài sản trí tuệ năm 2024, phần còn lại sau khi trích 50% về Trường "
         "theo Quy chế chi tiêu nội bộ vốn không nêu rõ phần của tác giả, hay 50% giảm dần xuống 20% theo Điều lệ Quỹ.")
    thay(doc, "Thứ hai, trong ba văn bản, Điều lệ Quỹ", "Thứ hai, trong các văn bản trên, Điều lệ Quỹ")
    thay(doc, "hiện hành: không đặt trần và cho phép", "hiện hành: vừa không đặt trần, vừa cho phép")
    # --- 2.2.2
    thay(doc, "có 2 nhân sự theo danh sách năm 2026. Không có nhân sự nào được ghi nhận là chuyên trách về sở hữu trí tuệ.",
         "có 2 nhân sự theo danh sách năm 2026, đều có trình độ thạc sĩ với chuyên ngành thông tin - thư viện và khoa "
         "học môi trường; chức danh trưởng phòng do Hiệu trưởng kiêm nhiệm. Không có nhân sự nào được ghi nhận là chuyên "
         "trách hoặc có chuyên ngành đào tạo về sở hữu trí tuệ.")
    p = tim(doc, "Như vậy, vấn đề của Nhà trường không phải là thiếu đầu mối")
    dat_chu(p, "Như vậy, vấn đề của Nhà trường không phải là thiếu đầu mối hay thiếu phân công. Điều 11 Quy chế quản trị "
               "tài sản trí tuệ năm 2024 đã giao Phòng Khoa học Công nghệ nhiệm vụ nhận diện, ghi nhận tài sản trí tuệ, xây "
               "dựng quy trình và biểu mẫu khai báo, lập hồ sơ theo dõi và xúc tiến thương mại hóa, đồng thời giao Bộ phận "
               "Pháp chế thực hiện thủ tục xác lập quyền. Khoảng cách nằm ở khâu thực thi: bộ hồ sơ thu thập được không có "
               "biểu mẫu khai báo hay hồ sơ theo dõi tài sản trí tuệ do Phòng lập, và trên thực tế không đơn vị nào trong "
               "bốn đơn vị nêu trên bao quát trọn vẹn chu trình từ nhận diện kết quả nghiên cứu, đánh giá khả năng bảo hộ, "
               "xác lập quyền đến khai thác và duy trì hiệu lực. Mỗi đơn vị nắm một đoạn, và giữa các đoạn không có cơ chế "
               "chuyển giao hồ sơ. Đây là đặc điểm tổ chức được trở lại tại Mục 2.5.3.")
    # --- 2.2.3
    thay(doc, "không có dòng chi nào cho phí nộp đơn và duy trì hiệu lực.",
         "không có dòng chi nào cho phí nộp đơn và duy trì hiệu lực. Điều 35 Quyết định 213 giao Phòng Khoa học Công nghệ "
         "nộp đơn và lệ phí, nhưng Điều 38 về các khoản chi không có mục riêng cho lệ phí đăng ký và phí duy trì hiệu lực; "
         "Điều 13 Quy chế quản trị tài sản trí tuệ năm 2024 chỉ đề cập lệ phí xác lập quyền như một khoản được trừ khi "
         "phân chia lợi ích, tức sau khi đã có doanh thu.")
    thay(doc, "Nhà trường không thiếu quy định, mà có ba văn bản quy định chồng lấn",
         "Nhà trường không thiếu quy định, mà có bốn văn bản quy định chồng lấn")
    o(T[4], 1, 2, "38 đề tài; 19 đề tài được Trường cấp kinh phí với tổng 424,75 triệu đồng; 16 đề tài chỉ quy đổi giờ "
                  "nghiên cứu; 3 đề tài tự tìm tài trợ")
    thay(doc, "Tổng kinh phí của kênh này trong năm năm là 374,25 triệu đồng, mức cao nhất cho một đề tài là 80 triệu "
              "đồng, trong khi 15 trên 37 đề tài không được cấp kinh phí bằng tiền mà chỉ được quy đổi giờ nghiên cứu "
              "khoa học.",
         "Tổng kinh phí Nhà trường cấp cho kênh này trong năm năm là 424,75 triệu đồng, mức cao nhất cho một đề tài là 80 "
         "triệu đồng, trong khi 19 trên 38 đề tài không được Nhà trường cấp kinh phí bằng tiền, gồm 16 đề tài chỉ được "
         "quy đổi giờ nghiên cứu khoa học và 3 đề tài tự tìm nguồn tài trợ.")
    thay(doc, "quy định mức thưởng không quá 50 triệu đồng cho đề tài có đăng ký sở hữu trí tuệ",
         "quy định mức kinh phí xét duyệt không quá 50 triệu đồng cho loại đề tài có đăng ký sở hữu trí tuệ")
    thay(doc, "phải có chứng nhận đăng ký thành công thì mới được thưởng.",
         "phải có chứng nhận đăng ký thành công thì đề tài mới được xếp vào loại này và văn bằng mới được quy đổi giờ.")
    thay(doc, "với tổng cộng gần mười tỷ đồng trong ba kênh", "với tổng cộng khoảng mười tỷ đồng trong ba kênh")
    # --- 2.3.1
    thay(doc, "Nhà trường phát sinh tài sản ở tám nhóm đối tượng quyền nhưng mới xác lập quyền ở bốn nhóm.",
         "Nhà trường phát sinh tài sản ở chín nhóm đối tượng quyền nhưng mới có hoạt động xác lập quyền ở bốn nhóm, "
         "trong đó chỉ ba nhóm đã được cấp văn bằng.")
    thay(doc, "Thứ hai, Điều 34 Quyết định 213 không liệt kê giải pháp hữu ích và sưu tập dữ liệu, trong khi đây lại là hai "
              "nhóm mà sản phẩm đề tài cấp cơ sở rơi vào nhiều nhất.",
         "Thứ hai, giải pháp hữu ích và sưu tập dữ liệu, hai nhóm mà sản phẩm đề tài cấp cơ sở rơi vào nhiều nhất, đều đã "
         "có trong phạm vi tài sản của quy chế nội bộ: giải pháp hữu ích được liệt kê tại Điều 34 Quyết định 213, sưu tập "
         "dữ liệu được bổ sung dưới dạng cơ sở dữ liệu tại Điều 3 Quy chế quản trị tài sản trí tuệ năm 2024. Khoảng trống "
         "vì vậy không nằm ở phạm vi văn bản mà ở khâu nhận diện: chưa sản phẩm nào thuộc hai nhóm này được xác lập quyền.")
    thay(doc, "các loại tài sản được liệt kê tại Điều 34 Quyết định số 213/QĐ-ĐHTĐ,",
         "các loại tài sản được liệt kê tại Điều 34 Quyết định 213 và Điều 3 Quy chế quản trị tài sản trí tuệ năm 2024,")
    thay(doc, "Luật Sở hữu trí tuệ, Quyết định số 213/QĐ-ĐHTĐ và các danh mục",
         "Luật Sở hữu trí tuệ, Quyết định số 213/QĐ-ĐHTĐ, Quy chế quản trị tài sản trí tuệ năm 2024 và các danh mục")
    b = T[5]
    o(b, 0, 2, "Quyết định 213 hoặc Quy chế 2024 có liệt kê")
    for h, v in ((5, "Có, tại Quy chế 2024"), (9, "Có"), (12, "Có, tại Quy chế 2024")):
        assert b.rows[h].cells[2].text.startswith("Không"), b.rows[h].cells[1].text
        o(b, h, 2, v)
    # --- 2.3.2
    b = T[6]
    assert "thương hiệu" in b.rows[13].cells[6].text
    o(b, 13, 6, "4 thương hiệu, 6 hợp tác, 2 nghiên cứu")
    thay(doc, "không thuộc phạm vi phân tích.",
         "không thuộc phạm vi phân tích. Hai sáng chế có trong danh mục theo dõi của Nhà trường nhưng danh mục chưa ghi "
         "số đơn.")
    thay(doc, "Nhóm thứ nhất gồm 5 tài sản gắn với thương hiệu", "Nhóm thứ nhất gồm 4 tài sản gắn với thương hiệu")
    thay(doc, "hoạt động nghiên cứu với 470 sản phẩm khoa học độc lập",
         f"hoạt động nghiên cứu với khoảng {B.doc_lap} sản phẩm khoa học độc lập")
    # --- 2.3.3
    thay(doc, "cột sản phẩm nghiệm thu của 37 đề tài cấp cơ sở", "cột sản phẩm nghiệm thu của 38 đề tài cấp cơ sở")
    b = T[7]
    for h in (1, 2, 3):
        assert b.rows[h].cells[0].text.strip() == "2021"
        o(b, h, 0, "2022")
    assert "Citrus" in b.rows[4].cells[2].text
    o(b, 4, 4, "Chưa nộp đơn cho công thức")
    assert b.rows[12].cells[1].text.startswith("11 đề tài")
    o(b, 12, 1, "11 đề tài trên 38")
    thay(doc, "sản phẩm mô hình thiết bị năm 2021 đã không còn", "sản phẩm mô hình thiết bị năm 2022 đã không còn")
    p = tim(doc, "Đối chiếu hai bảng cho thấy một điểm quan trọng")
    dat_chu(p, "Đối chiếu hai bảng cho thấy một điểm quan trọng: ngoài trường hợp lá Quế hoa, chỉ có một điểm giao nhau "
               "có thể nhận diện giữa Bảng 2.8 và Bảng 2.7. Đề tài chiết xuất tinh dầu Citrus grandis, tức tinh dầu "
               "bưởi, năm 2023 có liên hệ rõ về tên gọi và bối cảnh hợp tác với kiểu dáng công nghiệp tinh dầu vỏ bưởi "
               "đào được cấp năm 2024, song hồ sơ hiện có chưa đủ để khẳng định kiểu dáng này hình thành từ đề tài. Kể "
               "cả khi liên hệ được xác nhận, văn bằng kiểu dáng chỉ bảo hộ hình dáng bên ngoài của sản phẩm; công thức "
               "và quy trình chiết xuất, là phần kết quả nghiên cứu thực chất, vẫn chưa được đăng ký. Hai danh mục do hai "
               "bộ phận khác nhau lập và không có trường dữ liệu liên kết với nhau. Đây là bằng chứng trực tiếp cho nhận "
               "định về hai luồng tài sản độc lập đã nêu tại Mục 2.2.2.")
    # --- 2.4.1
    p = tim(doc, "Hình thức khai thác phổ biến nhất tại Nhà trường")
    dat_chu(p, "Hình thức khai thác phổ biến nhất tại Nhà trường là sử dụng tài liệu giảng dạy trong đào tạo. Trong 87 "
               "giáo trình và tài liệu được biên soạn giai đoạn 2021 - 2025, toàn bộ 61 tài liệu của giai đoạn 2022 - 2025 "
               "được ghi nhận số tín chỉ, tức gắn với một học phần cụ thể trong chương trình đào tạo. 26 tài liệu năm 2021 "
               "không có thông tin này do biểu mẫu danh mục năm 2021 chưa có trường số tín chỉ, chứ không có nghĩa là "
               "không được sử dụng. Đây là bằng chứng cho thấy tài liệu biên soạn được đưa vào sử dụng thực tế chứ không "
               "dừng ở sản phẩm nghiệm thu.")
    # --- 2.5.1
    thay(doc, "việc ban hành quy định từ năm 2021 là sớm.",
         "việc ban hành quy định từ năm 2021 là sớm. Năm 2024, Nhà trường tiếp tục ban hành Quy chế quản trị tài sản trí "
         "tuệ riêng, với phạm vi tài sản rộng, nguyên tắc công bố và bảo mật, cùng sự phân công đầu mối rõ ràng.")
    # --- 2.5.2
    thay(doc, "Độ phủ khai báo sản phẩm trong hồ sơ nghiệm thu chỉ đạt 26,2%.",
         "Độ phủ khai báo sản phẩm trong hồ sơ nghiệm thu chỉ khoảng 29,9%.")
    thay(doc, "Trong 37 đề tài cấp cơ sở, 11 đề tài", "Trong 38 đề tài cấp cơ sở, 11 đề tài")
    thay(doc, "Hoạt động nghiên cứu với khoảng 470 sản phẩm độc lập",
         f"Hoạt động nghiên cứu với khoảng {B.doc_lap} sản phẩm độc lập")
    thay(doc, "cho thấy hai danh mục không có điểm giao nhau, ngoại trừ một trường hợp duy nhất của năm 2025.",
         "cho thấy hai danh mục gần như không có điểm giao nhau: ngoài trường hợp lá Quế hoa năm 2025, chỉ có một liên "
         "hệ chưa được xác nhận bằng hồ sơ giữa đề tài tinh dầu bưởi và một kiểu dáng công nghiệp.")
    p = tim(doc, "Thứ sáu, khâu kiểm tra và giám sát")
    dat_chu(p, "Thứ sáu, khâu kiểm tra và giám sát chưa gắn với sản phẩm. Trong 38 đề tài, có 1 đề tài xếp loại Xuất sắc, "
               "19 đề tài xếp loại Tốt, 1 đề tài xếp loại Khá và 17 đề tài xếp loại Đạt. Bảy đề tài chỉ có sản phẩm là báo "
               "cáo tổng kết, nhưng 5 trong số đó vẫn được xếp loại Tốt, cho thấy kết quả đánh giá chưa phân biệt rõ mức độ "
               "tạo ra sản phẩm có thể ứng dụng hoặc bảo hộ.")
    # --- 2.5.3
    thay(doc, "Các hạn chế nêu trên bắt nguồn từ sáu nguyên nhân,",
         "Các hạn chế nêu trên bắt nguồn từ hai nhóm nguyên nhân khách quan và chủ quan,")
    thay(doc, "Về thể chế, ba văn bản nội bộ quy định ba công thức chia lợi ích khác nhau cho cùng một loại tài sản, với "
              "ba mức trần khác nhau và ba phạm vi áp dụng không rõ ranh giới.",
         "Về thể chế, bốn văn bản nội bộ chứa năm quy định chia lợi ích khác nhau cho cùng một loại tài sản, với các mức "
         "trần khác nhau và phạm vi áp dụng không rõ ranh giới.")
    thay(doc, "Bên cạnh đó, Điều 34 Quyết định 213 không liệt kê giải pháp hữu ích và sưu tập dữ liệu, trong khi đây lại là "
              "hai nhóm mà sản phẩm đề tài rơi vào nhiều nhất, dẫn đến các sản phẩm này không được nhận diện là đối tượng "
              "cần xác lập quyền.",
         "Bên cạnh đó, Quy chế quản trị tài sản trí tuệ năm 2024 không dẫn chiếu và không thay thế Chương VI Quyết định "
         "213, nên hai văn bản cùng tồn tại với phạm vi tài sản, phân công đầu mối và cơ chế phân chia lợi ích khác nhau.")
    thay(doc, "quy trình tồn tại trên văn bản nhưng không được kích hoạt.",
         "quy trình tồn tại trên văn bản nhưng không được kích hoạt. Điều 10 Quy chế quản trị tài sản trí tuệ năm 2024 tiếp "
         "tục đặt trách nhiệm này lên tác giả khi yêu cầu tác giả tự xác định tài sản có thể bảo hộ và xin ý kiến Phòng "
         "Khoa học Công nghệ trước khi bộc lộ công khai.")
    thay(doc, "Về tổ chức, chức năng quản lý quyền sở hữu trí tuệ bị chia cho bốn đơn vị mà không đơn vị nào bao quát trọn "
              "chu trình.",
         "Về tổ chức, Điều 11 Quy chế quản trị tài sản trí tuệ năm 2024 giao Phòng Khoa học Công nghệ vai trò bao quát chu "
         "trình, nhưng trên thực tế chức năng quản lý quyền sở hữu trí tuệ bị chia cho bốn đơn vị mà không đơn vị nào thực "
         "hiện trọn chu trình.")
    thay(doc, "chỉ đạt 26,2% càng thu hẹp", "chỉ khoảng 29,9% càng thu hẹp")
    p = tim(doc, "Về động lực, cơ chế khuyến khích")
    dat_chu(p, "Về động lực, cơ chế khuyến khích tồn tại trên văn bản với định mức quy đổi từ 180 đến 600 giờ nghiên cứu "
               "cho mỗi văn bằng, nhưng văn bằng không có khoản thưởng bằng tiền như bài báo quốc tế, vốn được thưởng từ "
               "10 đến 20 triệu đồng một bài. Quy chế đặt điều kiện phải có chứng nhận đăng ký thành công mới được ghi "
               "nhận, trong khi thủ tục xác lập quyền sở hữu công nghiệp thường kéo dài từ hai đến ba năm. Khoảng cách "
               "thời gian giữa nỗ lực và phần thưởng lớn hơn nhiều so với công bố khoa học, vốn được ghi nhận ngay khi "
               "bài được đăng, như phân tích tại Hình 2.11.")
    thay(doc, "Sáu nguyên nhân trên không độc lập với nhau.", "Bảy nguyên nhân chủ quan trên không độc lập với nhau.")


def them_quy_che_2024(doc):
    neo = tim(doc, "Văn bản đầu tiên điều chỉnh hoạt động sở hữu trí tuệ")
    p = doan_moi(neo._p, neo,
                 "Năm 2024, Nhà trường ban hành Quy chế quản trị tài sản trí tuệ, văn bản chuyên biệt đầu tiên về lĩnh vực "
                 "này. Bản quy chế được cung cấp cho đề tài do Phó Hiệu trưởng ký nhưng chưa ghi số và ngày ban hành; theo "
                 "Điều 17, quy chế có hiệu lực kể từ ngày ký. Quy chế có bốn nội dung chính. Điều 3 xác định phạm vi tài "
                 "sản rất rộng, bao gồm cả thông tin kỹ thuật có thể được cấp bằng độc quyền sáng chế hoặc giải pháp hữu "
                 "ích, cơ sở dữ liệu, phần mềm, giáo trình điện tử, bí quyết và tên miền, đồng thời quy định Nhà trường sở "
                 "hữu tài sản được tạo ra chủ yếu từ nguồn lực của Trường. Điều 10 quy định nguyên tắc công bố và bảo mật, "
                 "yêu cầu tác giả xin ý kiến Phòng Khoa học Công nghệ trước khi bộc lộ công khai tài sản có thể bảo hộ. Điều "
                 "11 giao Phòng Khoa học Công nghệ quản lý, nhận diện, lập hồ sơ theo dõi và xúc tiến thương mại hóa tài sản "
                 "trí tuệ, giao Bộ phận Pháp chế thực hiện thủ tục xác lập quyền. Điều 13 quy định lợi ích được phân chia "
                 "theo thỏa thuận; nếu không có thỏa thuận, Hiệu trưởng quyết định tỷ lệ sau khi trừ các chi phí. Quy chế "
                 "năm 2024 không dẫn chiếu và không thay thế Chương VI Quyết định 213, nên hai văn bản cùng tồn tại.")
    for r in p.runs:
        r.bold = False


def them_dan_nhap_25(doc):
    neo = tim(doc, "2.5. Đánh giá chung về công tác quản lý quyền sở hữu trí tuệ")
    mau = tim(doc, "Văn bản đầu tiên điều chỉnh hoạt động sở hữu trí tuệ")
    p = doan_moi(neo._p, mau,
                 "Trước khi đi vào từng nhóm kết quả và hạn chế, đề tài đối chiếu kết quả thực hiện với Kế hoạch hoạt động "
                 "khoa học công nghệ giai đoạn 2024 - 2028, tầm nhìn 2035, ban hành kèm Kế hoạch số 07/KH-ĐHTĐ ngày 01 "
                 "tháng 7 năm 2024. Đây là thước đo khách quan vì do chính Nhà trường đặt ra, và 2024 - 2025 là hai năm đầu "
                 "của kế hoạch.")
    p.paragraph_format.keep_with_next = False


def them_nguyen_nhan(doc):
    mo = tim(doc, "Các hạn chế nêu trên bắt nguồn từ hai nhóm nguyên nhân")
    mau = mo
    doan = [
        ("a) Nguyên nhân khách quan", True),
        ("Thứ nhất, khung pháp luật điều chỉnh sở hữu trí tuệ, khoa học công nghệ và giáo dục đại học thay đổi dồn dập "
         "trong giai đoạn 2025 - 2026. Chỉ trong khoảng một năm, Luật Khoa học, công nghệ và đổi mới sáng tạo số "
         "93/2025/QH15 có hiệu lực từ ngày 01 tháng 10 năm 2025, Luật Giáo dục đại học số 125/2025/QH15 được ban hành "
         "ngày 10 tháng 12 năm 2025, Luật số 131/2025/QH15 sửa đổi, bổ sung Luật Sở hữu trí tuệ có hiệu lực từ ngày 01 "
         "tháng 4 năm 2026, tiếp đó là Kết luận số 51-KL/TW ngày 17 tháng 6 năm 2026 của Bộ Chính trị và Quyết định số "
         "1624/QĐ-TTg ngày 21 tháng 8 năm 2026 sửa đổi Chiến lược sở hữu trí tuệ đến năm 2030. Quy chế nội bộ ban hành "
         "năm 2021 được xây dựng trên khung pháp luật trước đây, nên việc một số quy định như mức trần thù lao tại Điều "
         "36 Quyết định 213 không còn tương thích một phần là hệ quả của tốc độ thay đổi thể chế.", False),
        ("Thứ hai, thủ tục xác lập quyền sở hữu công nghiệp có đặc thù về thời gian và chi phí. Quá trình thẩm định đơn "
         "sáng chế và giải pháp hữu ích thường kéo dài từ hai đến ba năm, kèm chi phí tra cứu, soạn đơn, lệ phí và phí "
         "duy trì hiệu lực hằng năm. Đối với một trường đại học tư thục, các khoản chi này phải được cân đối từ nguồn thu "
         "của chính Nhà trường, không có nguồn ngân sách chi thường xuyên cho hoạt động khoa học công nghệ.", False),
        (f"Thứ ba, cơ cấu ngành đào tạo của Nhà trường chủ yếu thuộc các lĩnh vực kinh tế, quản lý, ngôn ngữ, giáo dục và "
         f"pháp luật, vốn chủ yếu tạo ra tác phẩm thuộc quyền tác giả và ít phát sinh đối tượng sở hữu công nghiệp. Dữ "
         f"liệu tại Hình 2.14 cho thấy {len(B.dt_duoc)} trên 11 sản phẩm đủ điều kiện xác lập quyền liên quan đến lĩnh "
         f"vực dược. Tiềm năng sở hữu công nghiệp vì vậy tập trung ở một lĩnh vực hẹp, khiến quy mô tài sản của Nhà "
         f"trường khó so sánh với các trường đại học kỹ thuật.", False),
        (f"Các nguyên nhân khách quan giải thích vì sao quy mô tài sản trí tuệ của Nhà trường còn nhỏ, nhưng không giải "
         f"thích được vì sao {len(B.du_dk_den_2024)} đề tài có sản phẩm đủ điều kiện trong giai đoạn 2021 - 2024 đều "
         f"không được nộp đơn. Phần lớn hạn chế bắt nguồn từ các nguyên nhân chủ quan dưới đây.", False),
        ("b) Nguyên nhân chủ quan", True),
        ("Nhóm nguyên nhân chủ quan gồm bảy nguyên nhân, xuất phát từ cách thiết kế và vận hành hệ thống quản lý của chính "
         "Nhà trường.", False),
    ]
    el = mo._p
    for text, la_muc in doan:
        p = doan_moi(el, mau, text, bold=True if la_muc else None, italic=True if la_muc else None,
                     giu_voi_sau=la_muc)
        el = p._p


def chen_hinh(doc):
    mau_tieu_de = tim(doc, "Bảng 2.2. Sản phẩm khoa học")
    mau_nguon = tim(doc, "Nguồn: Tổng hợp của nhóm nghiên cứu từ các danh mục thống kê")
    mau_than = tim(doc, "Số liệu tại Bảng 2.2 cho thấy hai xu hướng")

    # Gỡ hình ảnh cũ: đoạn tiêu đề Hình 2.1, đoạn ảnh, đoạn nguồn
    cu = tim(doc, "Hình 2.1. Mức độ chuyển hóa từ sản phẩm khoa học")
    anh = cu._p.getnext()
    nguon_cu = anh.getnext()
    assert "<w:drawing" in etree.tostring(anh, encoding="unicode")
    assert "".join(nguon_cu.itertext()).startswith("Nguồn: Tổng hợp của nhóm nghiên cứu. Số 470")
    for x in (cu._p, anh, nguon_cu):
        x.getparent().remove(x)

    for n, h in enumerate(B.HINH, 1):
        neo = tim(doc, h["sau_doan"])
        tieu_de = doan_moi(neo._p, mau_tieu_de, f"Hình 2.{h['so']}. {h['tieu_de']}", bold=True, giu_voi_sau=True)
        p_bd = doan_moi(tieu_de._p, mau_tieu_de, "", giu_voi_sau=True)
        for r in p_bd._p.findall(qn("w:r")):
            p_bd._p.remove(r)
        pf = p_bd.paragraph_format
        pf.space_before = 0
        pf.space_after = 0
        pf.line_spacing = 1.0
        p_bd._p.append(nhung_bieu_do(doc, h, n))
        nguon = doan_moi(p_bd._p, mau_nguon, h["nguon"], italic=True)
        el = nguon._p
        for doan in h["binh_luan"]:
            el = doan_moi(el, mau_than, doan)._p


def kiem_tra(doc):
    loi = []
    for p in doc.paragraphs:
        t = p.text
        if "—" in t or "–" in t:
            loi.append(("gạch dài", t[:80]))
        if re.search(r"\b(tôi|chúng tôi)\b", t, flags=re.I):
            loi.append(("ngôi thứ nhất", t[:80]))
    for m in re.finditer(r"Hình 2\.(\d+)", "\n".join(p.text for p in doc.paragraphs)):
        if int(m.group(1)) > len(B.HINH):
            loi.append(("tham chiếu hình", m.group(0)))
    return loi


def va_workbook(path):
    """Áp dụng chuẩn hóa nhãn dữ liệu cho các biểu đồ trong workbook tổng hợp."""
    import shutil
    import zipfile
    tam = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/charts/chart"):
                data = chuan_hoa_nhan(data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tam, path)


def main():
    dung_workbook()
    va_workbook(RA_XLSX)
    doc = docx.Document(VAO)
    sua_so_lieu(doc)
    them_quy_che_2024(doc)
    them_dan_nhap_25(doc)
    them_nguyen_nhan(doc)
    chen_hinh(doc)
    dat_cap_de_muc(doc)
    loi = kiem_tra(doc)
    doc.core_properties.title = "Chương 2. Thực trạng quản lý quyền sở hữu trí tuệ tại Trường Đại học Thành Đô"
    doc.save(RA_DOCX)
    print("Đã ghi:", RA_DOCX)
    print("Đã ghi:", RA_XLSX)
    for x in loi:
        print("CẢNH BÁO:", x)


if __name__ == "__main__":
    main()
