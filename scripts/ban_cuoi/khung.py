# -*- coding: utf-8 -*-
"""Khung dựng văn bản Word dùng chung cho ba chương bản cuối.

Mọi chương dùng cùng định dạng trang, kiểu chữ và kiểu bảng lấy từ
"Chuong_2_Thuc_trang_hoan_chinh mới.docx" (thông qua lớp VanBan của build_gon.py),
nên ba tệp đồng nhất khi ghép vào báo cáo tổng kết.
"""
import copy
import os
import re
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, os.path.join(GOC, "scripts", "chuong2"))
import build_gon as BG  # noqa: E402

THU_MUC_RA = os.path.join(GOC, "Ban_cuoi")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def tach_markup(text):
    """Tách chuỗi có **đậm** và *nghiêng* thành danh sách (đoạn chữ, đậm, nghiêng)."""
    kq = []
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)", text):
        if m.group(1) is not None:
            kq.append((m.group(1), True, False))
        elif m.group(2) is not None:
            kq.append((m.group(2), False, True))
        else:
            kq.append((m.group(3), False, False))
    return kq


class VanBanChung(BG.VanBan):
    def doan_md(self, loai, text, bold=None, italic=None, giu=False):
        p = self.doan(loai, "", bold=bold, italic=italic, giu=giu)
        r0 = p.runs[0]._r
        mau = copy.deepcopy(r0)
        r0.getparent().remove(r0)
        for chu, dam, nghieng in tach_markup(text):
            r = copy.deepcopy(mau)
            for t in r.findall(qn("w:t")):
                r.remove(t)
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = etree.SubElement(r, qn("w:rPr"))
                r.remove(rpr)
                r.insert(0, rpr)
            if dam:
                for x in rpr.findall(qn("w:b")):
                    rpr.remove(x)
                etree.SubElement(rpr, qn("w:b"))
            if nghieng:
                for x in rpr.findall(qn("w:i")):
                    rpr.remove(x)
                etree.SubElement(rpr, qn("w:i"))
            t = etree.SubElement(r, qn("w:t"))
            t.text = chu
            t.set(XML_SPACE, "preserve")
            p._p.append(r)
        return p

    def than_md(self, *ds):
        for t in ds:
            self.doan_md("than", t)

    def bang(self, tieu_de, cot, dong, nguon, rong, can=None, dong_tong=False, hang_dam=(), tien_to="2"):
        """Như VanBan.bang nhưng ô có thể nhiều đoạn (ngăn bằng \\n), có markup **đậm**,
        và các hàng trong hang_dam được tô như hàng tiêu đề."""
        self.so_bang += 1
        self.doan("tieu_de", f"Bảng {tien_to + '.' if tien_to else ''}{self.so_bang}. {tieu_de}", bold=True, giu=True)
        tbl = etree.SubElement(etree.Element("x"), qn("w:tbl"))
        tbl.append(copy.deepcopy(self.tblPr))
        grid = etree.SubElement(tbl, qn("w:tblGrid"))
        tong = sum(rong)
        tw = [int(9070 * r / tong) for r in rong]
        for w in tw:
            etree.SubElement(grid, qn("w:gridCol")).set(qn("w:w"), str(w))
        can = can or (["left"] + ["center"] * (len(cot) - 1))
        for i, hang in enumerate([cot] + dong):
            tr = etree.SubElement(tbl, qn("w:tr"))
            trpr = etree.SubElement(tr, qn("w:trPr"))
            if i == 0:
                etree.SubElement(trpr, qn("w:tblHeader"))
                etree.SubElement(trpr, qn("w:cantSplit"))
            dau = i == 0 or (i - 1) in hang_dam
            cuoi = dong_tong and i == len(dong)
            for j, v in enumerate(hang):
                tc = etree.SubElement(tr, qn("w:tc"))
                tcpr = etree.SubElement(tc, qn("w:tcPr"))
                x = etree.SubElement(tcpr, qn("w:tcW"))
                x.set(qn("w:w"), str(tw[j]))
                x.set(qn("w:type"), "dxa")
                if dau:
                    shd = etree.SubElement(tcpr, qn("w:shd"))
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "DCE8F7")
                etree.SubElement(tcpr, qn("w:vAlign")).set(qn("w:val"), "center")
                for dong_o in str(v).split("\n"):
                    p = etree.SubElement(tc, qn("w:p"))
                    ppr = etree.SubElement(p, qn("w:pPr"))
                    sp_ = etree.SubElement(ppr, qn("w:spacing"))
                    sp_.set(qn("w:before"), "20")
                    sp_.set(qn("w:after"), "20")
                    sp_.set(qn("w:line"), "252")
                    sp_.set(qn("w:lineRule"), "auto")
                    ind = etree.SubElement(ppr, qn("w:ind"))
                    ind.set(qn("w:firstLine"), "0")
                    etree.SubElement(ppr, qn("w:jc")).set(qn("w:val"), "center" if dau else can[j])
                    for chu, dam, nghieng in tach_markup(dong_o) or [("", False, False)]:
                        r = etree.SubElement(p, qn("w:r"))
                        rpr = etree.SubElement(r, qn("w:rPr"))
                        if dau or cuoi or dam:
                            etree.SubElement(rpr, qn("w:b"))
                        if nghieng:
                            etree.SubElement(rpr, qn("w:i"))
                        sz = etree.SubElement(rpr, qn("w:sz"))
                        sz.set(qn("w:val"), "24")
                        t = etree.SubElement(r, qn("w:t"))
                        t.text = chu
                        t.set(XML_SPACE, "preserve")
        self._them(tbl)
        if nguon:
            self.doan("nguon", nguon, italic=True)
        return self.so_bang


def dat_cap_de_muc(doc, so_chuong, tu=0):
    """Gắn cấp đề mục cho mục lục tự động: tên chương cấp 0, x.y cấp 1, x.y.z cấp 2, tiểu kết cấp 1.
    tu: chỉ số đoạn bắt đầu xét (bỏ qua phần mục lục ở đầu báo cáo gộp)."""
    c = str(so_chuong)
    for p in doc.paragraphs[tu:]:
        t = p.text.strip()
        lvl = None
        if t == f"CHƯƠNG {c}":
            lvl = 0
        elif re.match(rf"^{c}\.\d\. ", t) or t.startswith("TIỂU KẾT"):
            lvl = 1
        elif re.match(rf"^{c}\.\d\.\d\. ", t):
            lvl = 2
        if lvl is None:
            continue
        ppr = p._p.get_or_add_pPr()
        for x in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(x)
        ol = etree.Element(qn("w:outlineLvl"))
        ol.set(qn("w:val"), str(lvl))
        rpr = ppr.find(qn("w:rPr"))
        if rpr is not None:
            rpr.addprevious(ol)
        else:
            ppr.append(ol)
        p.paragraph_format.keep_with_next = True


def kiem_tra(doc):
    loi = []
    tat_ca = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    for t in tat_ca:
        if "—" in t or "–" in t:
            loi.append(("gạch dài", t[:80]))
        if re.search(r"\b(tôi|chúng tôi)\b", t, flags=re.I):
            loi.append(("ngôi thứ nhất", t[:80]))
        if re.search(r"\([^)]*\b(Living Lab|Scopus|Web of Science)\b[^)]*\)", t):
            loi.append(("ngoặc đơn giải thích tiếng Anh", t[:80]))
        if "**" in t:
            loi.append(("còn dấu markup", t[:80]))
    return loi


def luu(v, ten, tieu_de, so_chuong):
    dat_cap_de_muc(v.doc, so_chuong)
    v.doc.core_properties.title = tieu_de
    os.makedirs(THU_MUC_RA, exist_ok=True)
    ra = os.path.join(THU_MUC_RA, ten)
    v.doc.save(ra)
    so_tu = sum(len(p.text.split()) for p in v.doc.paragraphs)
    print("Đã ghi:", ra, f"({so_tu} từ ngoài bảng, {v.so_bang} bảng, {v.so_hinh} hình)")
    for x in kiem_tra(v.doc):
        print("CẢNH BÁO:", x)
    return ra


# ---------------------------------------------------------------------------
# Sửa đoạn tại chỗ sau khi dựng (dùng cho các lượt chỉnh sửa có bảng giải trình)
# ---------------------------------------------------------------------------
def _dat_chu(p, text):
    """Ghi lại nội dung đoạn p theo markup **đậm**, *nghiêng*, giữ định dạng của run đầu."""
    runs = p._p.findall(qn("w:r"))
    mau = copy.deepcopy(runs[0]) if runs else etree.Element(qn("w:r"))
    co_chu = [r for r in runs if "".join(t.text or "" for t in r.findall(qn("w:t"))).strip()]
    def _dam(r):
        rpr = r.find(qn("w:rPr"))
        b = rpr.find(qn("w:b")) if rpr is not None else None
        return b is not None and b.get(qn("w:val")) not in ("0", "false")
    ca_doan_dam = bool(co_chu) and all(_dam(r) for r in co_chu)
    for r in runs:
        r.getparent().remove(r)
    for x in mau.findall(qn("w:t")) + mau.findall(qn("w:br")):
        mau.remove(x)
    rpr0 = mau.find(qn("w:rPr"))
    if rpr0 is not None and not ca_doan_dam:
        for x in rpr0.findall(qn("w:b")) + rpr0.findall(qn("w:bCs")) + rpr0.findall(qn("w:i")):
            rpr0.remove(x)
    for chu, dam, nghieng in tach_markup(text):
        r = copy.deepcopy(mau)
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("w:rPr"))
            r.insert(0, rpr)
        if dam and rpr.find(qn("w:b")) is None:
            etree.SubElement(rpr, qn("w:b"))
        if nghieng and rpr.find(qn("w:i")) is None:
            etree.SubElement(rpr, qn("w:i"))
        t = etree.SubElement(r, qn("w:t"))
        t.text = chu
        t.set(XML_SPACE, "preserve")
        p._p.append(r)
    del rpr0


def thay_doan(doc, bat_dau, moi, tu=0, giai_trinh=None, vi_tri="", can_cu=""):
    """Thay đoạn duy nhất bắt đầu bằng bat_dau (tính từ đoạn thứ tu) bằng moi.
    moi là chuỗi hoặc danh sách chuỗi; chuỗi rỗng xóa đoạn; các chuỗi sau được chèn thành đoạn mới ngay sau.
    Nếu có danh sách giai_trinh, ghi (vị trí, nội dung cũ, nội dung mới, căn cứ)."""
    ps = doc.paragraphs[tu:]
    trung = [p for p in ps if p.text.strip().startswith(bat_dau)]
    assert len(trung) == 1, f"{len(trung)} đoạn bắt đầu bằng: {bat_dau[:70]}"
    p = trung[0]
    cu = p.text.strip()
    moi = [moi] if isinstance(moi, str) else list(moi)
    if moi == [""]:
        p._p.getparent().remove(p._p)
    else:
        neo = p._p
        _dat_chu(p, moi[0])
        for t in moi[1:]:
            el = copy.deepcopy(p._p)
            neo.addnext(el)
            neo = el
            _dat_chu(docx.text.paragraph.Paragraph(el, p._parent), t)
    if giai_trinh is not None:
        giai_trinh.append((vi_tri, cu, "\n".join(re.sub(r"\*", "", x) for x in moi) if moi != [""] else "(lược bỏ)",
                           can_cu))
    return p
