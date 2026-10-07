# -*- coding: utf-8 -*-
"""Gộp ba chương thành một báo cáo tổng kết hoàn chỉnh.

    python3 scripts/ban_cuoi/bao_cao.py

Đầu ra: Ban_cuoi/Bao_cao_tong_ket_de_tai.docx và Ban_cuoi/Du_lieu_bieu_do.xlsx.

Cấu trúc: trang bìa, mục lục, danh mục bảng, danh mục hình (đánh số trang La Mã),
Mở đầu, Chương 1 - 3, Kết luận và kiến nghị, Tài liệu tham khảo (đánh số trang Ả Rập
từ 1). Mục lục và hai danh mục là trường TOC của Word; để tệp mở ra đã có sẵn nội
dung, số trang được lấy từ một lần kết xuất PDF bằng LibreOffice rồi ghi vào phần kết
quả của trường. Word cập nhật lại khi mở tệp (settings/updateFields).
"""
import copy
import os
import re
import subprocess
import sys
import tempfile

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import khung  # noqa: E402
import chuong1  # noqa: E402
import chuong3  # noqa: E402
import mo_dau_ket_luan as MK  # noqa: E402
import tai_lieu as TL  # noqa: E402

BG = khung.BG
B, BD = BG.B, BG.BD
GOC = khung.GOC
RA_DOCX = os.path.join(khung.THU_MUC_RA, "Bao_cao_tong_ket_de_tai.docx")
RA_XLSX = os.path.join(khung.THU_MUC_RA, "Du_lieu_bieu_do.xlsx")
TEN_DE_TAI = "Nghiên cứu giải pháp nâng cao hiệu quả quản lý quyền sở hữu trí tuệ tại Trường Đại học Thành Đô"
KIEU_BANG, KIEU_HINH = "Chu thich bang", "Chu thich hinh"
RONG_CHU = 9070  # 16 cm vùng chữ, đơn vị twip


# ---------------------------------------------------------------------------
# Tiện ích đoạn văn
# ---------------------------------------------------------------------------
def dong(v, text, co=13, dam=False, nghieng=False, truoc=0, sau=0, can=WD_ALIGN_PARAGRAPH.CENTER):
    p = v.doan("chuong1", text, bold=dam, italic=nghieng)
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.alignment = Pt(truoc), Pt(sau), can
    for r in p.runs:
        r.font.size = Pt(co)
    return p


def dat_outline(p, lvl):
    ppr = p._p.get_or_add_pPr()
    for x in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(x)
    ol = etree.Element(qn("w:outlineLvl"))
    ol.set(qn("w:val"), str(lvl))
    rpr = ppr.find(qn("w:rPr"))
    rpr.addprevious(ol) if rpr is not None else ppr.append(ol)


THU_TU_SECT = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
               "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg",
               "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings"]


def chen_sect(sect, ten, **thuoc_tinh):
    """Chèn phần tử con vào sectPr đúng thứ tự lược đồ."""
    for x in sect.findall(qn("w:" + ten)):
        sect.remove(x)
    el = etree.Element(qn("w:" + ten))
    for k, gt in thuoc_tinh.items():
        el.set(qn("w:" + k), gt)
    vt = THU_TU_SECT.index(ten)
    sau = [c for c in sect if etree.QName(c).localname in THU_TU_SECT[vt + 1:]]
    sau[0].addprevious(el) if sau else sect.append(el)
    return el


def run_truong(p_el, loai, instr=None):
    r = etree.SubElement(p_el, qn("w:r"))
    if loai == "instr":
        t = etree.SubElement(r, qn("w:instrText"))
        t.set(khung.XML_SPACE, "preserve")
        t.text = instr
    else:
        etree.SubElement(r, qn("w:fldChar")).set(qn("w:fldCharType"), loai)
    return r


def run_chu(p_el, text, dam=False, co=None):
    r = etree.SubElement(p_el, qn("w:r"))
    rpr = etree.SubElement(r, qn("w:rPr"))
    if dam:
        etree.SubElement(rpr, qn("w:b"))
    if co:
        etree.SubElement(rpr, qn("w:sz")).set(qn("w:val"), str(co * 2))
    t = etree.SubElement(r, qn("w:t"))
    t.set(khung.XML_SPACE, "preserve")
    t.text = text
    return r


def doan_muc_luc(v, instr, muc, dam_cap0=True):
    """Chèn trường TOC. muc: danh sách (cấp, chữ, trang) làm kết quả sẵn có; rỗng thì ghi lời nhắc."""
    muc = muc or [(0, "Nhấn chuột phải, chọn Update Field để cập nhật.", "")]
    for i, (cap, chu, trang) in enumerate(muc):
        p = v.doan("than", "")
        for r in p._p.findall(qn("w:r")):
            p._p.remove(r)
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(14 * cap)
        pf.space_before, pf.space_after = Pt(0), Pt(3)
        pf.line_spacing = 1.15
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ppr = p._p.get_or_add_pPr()
        tabs = etree.Element(qn("w:tabs"))
        tab = etree.SubElement(tabs, qn("w:tab"))
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(RONG_CHU))
        ppr.insert(0, tabs)
        if i == 0:
            run_truong(p._p, "begin")
            run_truong(p._p, "instr", instr)
            run_truong(p._p, "separate")
        run_chu(p._p, chu, dam=dam_cap0 and cap == 0)
        if trang != "":
            etree.SubElement(etree.SubElement(p._p, qn("w:r")), qn("w:tab"))
            run_chu(p._p, str(trang), dam=dam_cap0 and cap == 0)
        if i == len(muc) - 1:
            run_truong(p._p, "end")


def tieu_de_phan(v, text, lvl=0, ngat_trang=True):
    p = v.doan("chuong1", text, bold=True)
    for r in p.runs:
        r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    if ngat_trang:
        p.paragraph_format.page_break_before = True
    if lvl is not None:
        dat_outline(p, lvl)
    return p


# ---------------------------------------------------------------------------
# Các phần của báo cáo
# ---------------------------------------------------------------------------
def trang_bia(v):
    dong(v, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", co=13)
    dong(v, "TRƯỜNG ĐẠI HỌC THÀNH ĐÔ", co=13, dam=True)
    dong(v, "--------------------", co=13, sau=90)
    dong(v, "BÁO CÁO TỔNG KẾT", co=18, dam=True, sau=6)
    dong(v, "ĐỀ TÀI KHOA HỌC VÀ CÔNG NGHỆ CẤP CƠ SỞ NĂM 2026", co=14, dam=True, sau=60)
    dong(v, "Tên đề tài:", co=13, nghieng=True, sau=6)
    dong(v, TEN_DE_TAI.upper(), co=17, dam=True, sau=110)
    dong(v, "Chủ nhiệm đề tài: Nguyễn Thị Tố Uyên", co=14, dam=True, can=WD_ALIGN_PARAGRAPH.LEFT).paragraph_format \
        .left_indent = Pt(100)
    dong(v, "Thành viên tham gia: Trần Đăng Bộ", co=14, can=WD_ALIGN_PARAGRAPH.LEFT, sau=150).paragraph_format \
        .left_indent = Pt(100)
    dong(v, "Hà Nội, năm 2026", co=14, dam=True)


def phan_dau(v, trang):
    trang_bia(v)
    tieu_de_phan(v, "MỤC LỤC", lvl=None)
    doan_muc_luc(v, 'TOC \\o "1-3" \\h \\z \\u', trang.get("muc_luc"))
    tieu_de_phan(v, "DANH MỤC BẢNG", lvl=None)
    doan_muc_luc(v, f'TOC \\h \\z \\t "{KIEU_BANG},1"', trang.get("bang"), dam_cap0=False)
    tieu_de_phan(v, "DANH MỤC HÌNH", lvl=None)
    doan_muc_luc(v, f'TOC \\h \\z \\t "{KIEU_HINH},1"', trang.get("hinh"), dam_cap0=False)
    # ngắt phần: phần đầu đánh số La Mã, trang bìa không số
    p_cuoi = v.body.findall(qn("w:p"))[-1]
    sect = copy.deepcopy(v.sect)
    chen_sect(sect, "pgNumType", fmt="lowerRoman")
    chen_sect(sect, "titlePg")
    p_cuoi.get_or_add_pPr().append(sect)


def phan_muc(v, tieu_de, cac_muc):
    tieu_de_phan(v, tieu_de)
    for ten, doan in cac_muc:
        dat_outline(v.doan("h1", ten), 1)
        v.than_md(*doan)


def tai_lieu_tham_khao(v, van_ban):
    tieu_de_phan(v, "TÀI LIỆU THAM KHẢO")
    dung = TL.duoc_trich(van_ban)
    for ma, ten in TL.NHOM:
        ds = [m for m in dung if m[1] == ma]
        if not ds:
            continue
        v.doan("h2", ten, bold=True)
        for apa in TL.danh_muc(ds):
            p = v.doan_md("than", apa)
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Pt(28), Pt(-28)
            pf.space_after = Pt(4)
    return dung


# ---------------------------------------------------------------------------
# Hậu xử lý
# ---------------------------------------------------------------------------
def them_kieu_chu_thich(doc):
    for ten in (KIEU_BANG, KIEU_HINH):
        if ten not in [s.name for s in doc.styles]:
            s = doc.styles.add_style(ten, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = doc.styles["Normal"]
            s.quick_style = False


def gan_kieu_chu_thich(doc):
    """Gắn kiểu cho dòng tên bảng, tên hình để trường TOC \\t lập danh mục."""
    for p in doc.paragraphs[dau_noi_dung(doc):]:
        t = p.text.strip()
        m = re.match(r"^(Bảng|Hình) \d\.\d+\. ", t)
        if not m or not p.runs or not p.runs[0].bold:
            continue
        p.style = doc.styles[KIEU_BANG if m.group(1) == "Bảng" else KIEU_HINH]


def chan_trang(doc):
    """Số trang giữa chân trang: phần đầu La Mã (bỏ trang bìa), phần chính Ả Rập từ 1."""
    def them_so(footer):
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_truong(p._p, "begin")
        run_truong(p._p, "instr", " PAGE ")
        run_truong(p._p, "separate")
        run_chu(p._p, "1", co=12)
        run_truong(p._p, "end")

    s1, s2 = doc.sections[0], doc.sections[1]
    them_so(s1.footer)
    s1.first_page_footer.is_linked_to_previous = False
    them_so(s2.footer)
    chen_sect(s2._sectPr, "pgNumType", fmt="decimal", start="1")


def bat_cap_nhat_truong(doc):
    st = doc.settings.element
    for x in st.findall(qn("w:updateFields")):
        st.remove(x)
    u = etree.Element(qn("w:updateFields"))
    u.set(qn("w:val"), "true")
    sau = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
           "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
           "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
           "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
    ke = [c for c in st if etree.QName(c).localname in sau]
    ke[0].addprevious(u) if ke else st.append(u)


def gop_ten_chuong(doc):
    """Gộp dòng "CHƯƠNG x" và các dòng tên chương thành một đoạn (ngắt dòng) để mục lục có tên chương đầy đủ."""
    ps = doc.paragraphs
    for i, p in enumerate(ps):
        if i < dau_noi_dung(doc) or not re.match(r"^CHƯƠNG \d$", p.text.strip()):
            continue
        p.paragraph_format.page_break_before = True
        for q in ps[i + 1:i + 5]:
            t = q.text.strip()
            if not t or t != t.upper() or re.match(r"^\d", t):
                break
            etree.SubElement(etree.SubElement(p._p, qn("w:r")), qn("w:br"))
            for r in q._p.findall(qn("w:r")):
                p._p.append(r)
            q._p.getparent().remove(q._p)


def dau_noi_dung(doc):
    """Chỉ số đoạn đầu tiên sau ngắt phần (hết phần mục lục, danh mục)."""
    return next(i for i, p in enumerate(doc.paragraphs)
                if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None) + 1


def hau_xu_ly(v):
    doc = v.doc
    for c in (1, 2, 3):
        khung.dat_cap_de_muc(doc, c, tu=dau_noi_dung(doc))
    gop_ten_chuong(doc)
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.match(r"^[ab]\) Nguyên nhân", t):
            p.paragraph_format.keep_with_next = True
    them_kieu_chu_thich(doc)
    gan_kieu_chu_thich(doc)
    chan_trang(doc)
    bat_cap_nhat_truong(doc)


# ---------------------------------------------------------------------------
# Dựng và lấy số trang
# ---------------------------------------------------------------------------
def dung_van_ban(trang):
    v = khung.VanBanChung()
    phan_dau(v, trang)
    phan_muc(v, "MỞ ĐẦU", MK.MO_DAU)
    v.so_bang = v.so_hinh = 0
    chuong1.noi_dung(v)
    v.so_bang = v.so_hinh = 0
    BG.noi_dung(v)
    h_kh = v.ket_qua["h_kh"]
    v.so_bang = v.so_hinh = 0
    chuong3.noi_dung(v, so_hinh_kh=h_kh)
    phan_muc(v, "KẾT LUẬN VÀ KIẾN NGHỊ", MK.KET_LUAN)
    van_ban = "\n".join([p.text for p in v.doc.paragraphs] +
                        [c.text for t in v.doc.tables for r in t.rows for c in r.cells])
    dung = tai_lieu_tham_khao(v, van_ban)
    hau_xu_ly(v)
    return v, van_ban, dung


def ket_xuat_pdf(duong_dan, thu_muc):
    subprocess.run(["soffice", "--headless", "-env:UserInstallation=file:///tmp/lo_prof_bc", "--convert-to", "pdf",
                    "--outdir", thu_muc, duong_dan], check=True, capture_output=True, timeout=600)
    return os.path.join(thu_muc, os.path.splitext(os.path.basename(duong_dan))[0] + ".pdf")


def lay_so_trang(v, pdf):
    import pymupdf
    trang_pdf = [re.sub(r"\s+", " ", pg.get_text()) for pg in pymupdf.open(pdf)]
    bat_dau = next(i for i, t in enumerate(trang_pdf) if t.strip().startswith("MỞ ĐẦU"))
    muc, bang, hinh = [], [], []
    vi_tri = bat_dau
    for p in v.doc.paragraphs[dau_noi_dung(v.doc):]:
        t = p.text.strip()
        if not t:
            continue
        ppr = p._p.pPr
        ol = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        kieu = p.style.name if p.style is not None else ""
        if ol is None and kieu not in (KIEU_BANG, KIEU_HINH):
            continue
        khoa = re.sub(r"\s+", " ", t)[:45]
        for i in range(vi_tri, len(trang_pdf)):
            if khoa in trang_pdf[i]:
                vi_tri = i
                break
        else:
            print("Không thấy trên PDF:", khoa)
        so = vi_tri - bat_dau + 1
        if ol is not None:
            dong_ = [x.strip() for x in t.split("\n")]
            ten = dong_[0] + (". " + " ".join(dong_[1:]) if len(dong_) > 1 else "")
            muc.append((int(ol.get(qn("w:val"))), ten, so))
        elif kieu == KIEU_BANG:
            bang.append((0, t, so))
        else:
            hinh.append((0, t, so))
    return {"muc_luc": muc, "bang": bang, "hinh": hinh}, len(trang_pdf) - bat_dau


def ghi_workbook(v):
    B.HINH[:] = v.hinh
    BD.RA_XLSX = RA_XLSX
    if BG.NHAT_KY_GON[0] not in BD.NHAT_KY:
        BD.NHAT_KY[:0] = BG.NHAT_KY_GON
    BD.dung_workbook()
    BD.va_workbook(RA_XLSX)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        v, _, _ = dung_van_ban({})
        nhap = os.path.join(tmp, "nhap.docx")
        v.doc.save(nhap)
        trang, so_trang = lay_so_trang(v, ket_xuat_pdf(nhap, tmp))
    for h in B.HINH:  # lần dựng thứ hai đánh số lại hình từ đầu
        h["id"] = h["id_cu"]
    v, van_ban, dung = dung_van_ban(trang)
    v.doc.core_properties.title = "Báo cáo tổng kết đề tài: " + TEN_DE_TAI
    v.doc.core_properties.author = "Nguyễn Thị Tố Uyên, Trần Đăng Bộ"
    v.doc.save(RA_DOCX)
    ghi_workbook(v)
    so_tu = sum(len(p.text.split()) for p in v.doc.paragraphs)
    print("Đã ghi:", RA_DOCX, f"({so_tu} từ ngoài bảng, khoảng {so_trang} trang nội dung, "
                             f"{len(trang['bang'])} bảng, {len(trang['hinh'])} hình, {len(dung)} tài liệu)")
    print("Đã ghi:", RA_XLSX)
    mo_coi = TL.trich_dan_mo_coi(van_ban)
    if mo_coi:
        print("CẢNH BÁO trích dẫn không có trong danh mục:", sorted(set(mo_coi)))
    for x in khung.kiem_tra(v.doc):
        print("CẢNH BÁO:", x)


if __name__ == "__main__":
    main()
