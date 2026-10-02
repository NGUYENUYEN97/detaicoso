# -*- coding: utf-8 -*-
"""Bài báo "Quản lý sở hữu trí tuệ tại trường đại học trước yêu cầu chính sách mới" theo khung của chủ nhiệm đề tài.

    python3 scripts/ban_cuoi/bai_bao_chinh_sach_moi.py

Đầu ra: Ban_cuoi/Bai_bao_Quan_ly_SHTT_yeu_cau_chinh_sach_moi.docx (bản thảo) và
        Ban_cuoi/To_khai_AI_Bai_bao_Quan_ly_SHTT.docx (tờ khai nộp kèm qua hệ thống OJS).
Thể lệ trình bày: Times New Roman 10,5; giãn dòng 11 pt; A4; lề trên, dưới 2,5 cm, trái, phải 2 cm; phần đầu bài
(tiêu đề, tóm tắt song ngữ) một cột, nội dung chính hai cột; bảng và hình rộng được đặt trong đoạn một cột.
Nội dung ở bai_bao_chinh_sach_moi_nd.py.
"""
import copy
import os
import re
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import khung  # noqa: E402
import so_do  # noqa: E402
import bai_bao_chinh_sach_moi_nd as N  # noqa: E402

RA_DOCX = os.path.join(khung.THU_MUC_RA, "Bai_bao_Quan_ly_SHTT_yeu_cau_chinh_sach_moi.docx")
RA_TO_KHAI = os.path.join(khung.THU_MUC_RA, "To_khai_AI_Bai_bao_Quan_ly_SHTT.docx")
CO_CHU, CO_BANG = 10.5, 9.5


# ---------------------------------------------------------------------------
# Tiện ích định dạng
# ---------------------------------------------------------------------------
def tao_tai_lieu():
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(CO_CHU)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(11)
    pf.space_before, pf.space_after = Pt(0), Pt(3)
    for z in d.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")  # mẫu mặc định của python-docx thiếu thuộc tính bắt buộc này
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.0)
    return d


def them_chu(p, text, co=None, dam=False, nghieng=False):
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)", text):
        chu = m.group(1) or m.group(2) or m.group(3)
        r = p.add_run(chu)
        r.bold = dam or m.group(1) is not None
        r.italic = nghieng or m.group(2) is not None
        if co:
            r.font.size = Pt(co)
    return p


def doan(d, text, can=WD_ALIGN_PARAGRAPH.JUSTIFY, thut=0.6, co=None, dam=False, nghieng=False, truoc=0, sau=3):
    p = d.add_paragraph()
    p.alignment = can
    pf = p.paragraph_format
    pf.first_line_indent = Cm(thut)
    pf.space_before, pf.space_after = Pt(truoc), Pt(sau)
    return them_chu(p, text, co, dam, nghieng)


def tieu_de_muc(d, text, cap=1):
    p = doan(d, text, can=WD_ALIGN_PARAGRAPH.LEFT, thut=0, dam=True, truoc=6 if cap == 1 else 4, sau=3)
    p.paragraph_format.keep_with_next = True
    return p


def ngat_phan(d, so_cot_phan_truoc):
    """Kết thúc phần hiện tại (ngắt liên tục). so_cot_phan_truoc: số cột của phần vừa kết thúc."""
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    sect = copy.deepcopy(d.sections[-1]._sectPr)
    for x in sect.findall(qn("w:cols")) + sect.findall(qn("w:type")):
        sect.remove(x)
    loai = etree.Element(qn("w:type"))
    loai.set(qn("w:val"), "continuous")
    sect.insert(0, loai) if sect.find(qn("w:headerReference")) is None else \
        sect.findall(qn("w:footerReference"))[-1].addnext(loai)
    cols = etree.SubElement(sect, qn("w:cols"))
    cols.set(qn("w:num"), str(so_cot_phan_truoc))
    cols.set(qn("w:space"), str(int(0.6 * 567)))
    # cols phải đứng trước docGrid theo lược đồ
    grid = sect.find(qn("w:docGrid"))
    if grid is not None:
        grid.addprevious(cols)
    p._p.get_or_add_pPr().append(sect)


def dat_cot_phan_cuoi(d, so_cot):
    sect = d.sections[-1]._sectPr
    for x in sect.findall(qn("w:cols")) + sect.findall(qn("w:type")):
        sect.remove(x)
    loai = etree.Element(qn("w:type"))
    loai.set(qn("w:val"), "continuous")
    sect.insert(0, loai)
    cols = etree.Element(qn("w:cols"))
    cols.set(qn("w:num"), str(so_cot))
    cols.set(qn("w:space"), str(int(0.6 * 567)))
    grid = sect.find(qn("w:docGrid"))
    grid.addprevious(cols) if grid is not None else sect.append(cols)


def to_nen(cell, mau):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = etree.SubElement(tcpr, qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau)


def bang(d, so, b):
    ten = f"Bảng {so}. {b['tieu_de']}" if isinstance(so, int) else b["tieu_de"]
    p = doan(d, f"**{ten}**", can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, truoc=4, sau=3)
    p.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(b["cot"]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tong = sum(b["rong"])
    rong = [Cm(17.0 * r / tong) for r in b["rong"]]

    def o(cell, text, dau=False):
        cell.text = ""
        pp = cell.paragraphs[0]
        pp.paragraph_format.first_line_indent = Cm(0)
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = Pt(10.5)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if dau else WD_ALIGN_PARAGRAPH.LEFT
        them_chu(pp, text, co=CO_BANG, dam=dau)

    for j, c in enumerate(b["cot"]):
        o(t.rows[0].cells[j], c, dau=True)
        to_nen(t.rows[0].cells[j], "DCE8F7")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    etree.SubElement(trpr, qn("w:tblHeader"))
    for hang in b["dong"]:
        cells = t.add_row().cells
        if all(x == "" for x in hang[1:]):
            gop = cells[0].merge(cells[-1])
            o(gop, hang[0])
            to_nen(gop, "F2F2F2")
            continue
        for j, x in enumerate(hang):
            o(cells[j], x)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            etree.SubElement(trpr, qn("w:cantSplit"))
        for j, w in enumerate(rong):
            if j < len(row.cells):
                row.cells[j].width = w
    if b["nguon"]:
        doan(d, b["nguon"], can=WD_ALIGN_PARAGRAPH.LEFT, thut=0, co=9.5, nghieng=True, truoc=2, sau=6)
    else:
        doan(d, "", thut=0, sau=4)


def hinh(d, so, ten, anh, nguon, rong_cm=16.0):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(anh, width=Cm(rong_cm))
    q = doan(d, f"**Hình {so}. {ten}**", can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, truoc=3, sau=2)
    q.paragraph_format.keep_with_next = True
    doan(d, nguon, can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, co=9.5, nghieng=True, sau=6)


def dem_tu(ds):
    return sum(len(re.sub(r"\*", "", x).split()) for x in ds)


# ---------------------------------------------------------------------------
# Dựng bản thảo
# ---------------------------------------------------------------------------
def dung_bai():
    d = tao_tai_lieu()
    dem = {}

    def phan(ten, ds):
        dem[ten] = dem.get(ten, 0) + dem_tu(ds)
        for x in ds:
            doan(d, x)

    # Phần đầu, một cột
    doan(d, N.TIEU_DE.upper(), can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, co=14, dam=True, sau=6)
    doan(d, N.TAC_GIA, can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, dam=True, sau=1)
    doan(d, N.DON_VI, can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, nghieng=True, sau=8)
    doan(d, "**Tóm tắt:** " + N.TOM_TAT, thut=0)
    doan(d, "**Từ khóa:** " + N.TU_KHOA, thut=0, sau=8)
    doan(d, N.TIEU_DE_EN.upper(), can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, co=12, dam=True, sau=4)
    doan(d, "**Abstract:** *" + N.ABSTRACT + "*", thut=0)
    doan(d, "**Keywords:** *" + N.KEYWORDS + "*", thut=0, sau=8)
    ngat_phan(d, 1)

    # Nội dung chính, hai cột
    tieu_de_muc(d, "1. ĐẶT VẤN ĐỀ")
    phan("Đặt vấn đề", N.DAT_VAN_DE)
    tieu_de_muc(d, "2. TỔNG QUAN NGHIÊN CỨU")
    tieu_de_muc(d, "2.1. Nghiên cứu về quản lý sở hữu trí tuệ trong trường đại học", 2)
    phan("Tổng quan", N.TONG_QUAN_21)
    tieu_de_muc(d, "2.2. Vấn đề tiếp tục nghiên cứu và hướng tiếp cận của bài báo", 2)
    phan("Tổng quan", N.TONG_QUAN_22)
    tieu_de_muc(d, "3. PHƯƠNG PHÁP NGHIÊN CỨU")
    phan("Phương pháp", N.PHUONG_PHAP)
    tieu_de_muc(d, "4. KẾT QUẢ NGHIÊN CỨU")
    tieu_de_muc(d, "4.1. Điểm thay đổi trong yêu cầu chính sách và hệ quả quản lý", 2)
    phan("Kết quả 4.1", N.KQ_41)
    ngat_phan(d, 2)
    bang(d, 1, N.BANG_1)
    ngat_phan(d, 1)
    phan("Kết quả 4.1", N.KQ_41_SAU)
    tieu_de_muc(d, "4.2. Cơ chế thực hiện qua năm nhóm giải pháp", 2)
    phan("Kết quả 4.2", N.KQ_42_MO)
    tieu_de_muc(d, "4.2.1. Quy chế xác định chủ thể quyết định và cơ chế lợi ích theo nguồn", 2)
    phan("Kết quả 4.2", N.GP1)
    tieu_de_muc(d, "4.2.2. Phân công theo chức năng thay vì mặc định lập bộ máy", 2)
    phan("Kết quả 4.2", N.GP2)
    tieu_de_muc(d, "4.2.3. Quy trình nộp đơn trước công bố, không chờ nghiệm thu", 2)
    phan("Kết quả 4.2", N.GP3)
    ngat_phan(d, 2)
    hinh(d, 1, "Quy trình phối hợp quản lý tài sản trí tuệ trong trường đại học",
         so_do.quy_trinh_bai_bao(), "Nguồn: Nhóm tác giả đề xuất trên cơ sở Chính phủ (2025), Văn phòng Quốc hội "
                                    "(2026) và World Intellectual Property Organization (2020).")
    ngat_phan(d, 1)
    phan("Kết quả 4.2", N.GP3_SAU)
    tieu_de_muc(d, "4.2.4. Nguồn chi cho đơn nộp trong quá trình nghiên cứu", 2)
    phan("Kết quả 4.2", N.GP4)
    tieu_de_muc(d, "4.2.5. Năng lực và văn hóa sở hữu trí tuệ theo vai trò", 2)
    phan("Kết quả 4.2", N.GP5)
    tieu_de_muc(d, "4.3. Lộ trình triển khai và chỉ số có định nghĩa", 2)
    phan("Kết quả 4.3", N.KQ_43)
    ngat_phan(d, 2)
    bang(d, 2, N.BANG_2)
    ngat_phan(d, 1)
    phan("Kết quả 4.3", N.KQ_43_SAU)
    tieu_de_muc(d, "5. BÀN LUẬN")
    phan("Bàn luận", N.BAN_LUAN)
    tieu_de_muc(d, "6. KẾT LUẬN")
    phan("Kết luận", N.KET_LUAN)
    tieu_de_muc(d, "TÀI LIỆU THAM KHẢO")
    for apa in N.TAI_LIEU:
        p = doan(d, apa, can=WD_ALIGN_PARAGRAPH.LEFT, thut=0, co=9.5, sau=2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
    dat_cot_phan_cuoi(d, 2)

    d.core_properties.title = N.TIEU_DE
    d.core_properties.author = N.TAC_GIA
    d.save(RA_DOCX)
    return d, dem


def dung_to_khai():
    d = tao_tai_lieu()
    doan(d, "TỜ KHAI MINH BẠCH SỬ DỤNG AI VÀ CAM KẾT DỮ LIỆU GỐC", can=WD_ALIGN_PARAGRAPH.CENTER, thut=0, co=13,
         dam=True, sau=8)
    doan(d, "Tên bài báo: " + N.TIEU_DE + ".", thut=0)
    doan(d, "Tác giả: " + N.TAC_GIA + ". Đơn vị: " + N.DON_VI + ".", thut=0, sau=6)
    doan(d, "**Phần I. Minh bạch sử dụng trí tuệ nhân tạo** (chọn một phương án)", thut=0)
    doan(d, "☐ Phương án A: Nhóm tác giả cam kết không sử dụng công cụ trí tuệ nhân tạo trong toàn bộ quá trình nghiên "
            "cứu, xử lý số liệu và viết bản thảo.", thut=0)
    doan(d, "☐ Phương án B: Nhóm tác giả xác nhận có sử dụng công cụ trí tuệ nhân tạo trong phạm vi cho phép, gồm nhuận "
            "ngôn ngữ, biên dịch câu chữ, gợi ý cấu trúc văn bản; không dùng trí tuệ nhân tạo để tạo nội dung khoa học, "
            "tạo dữ liệu giả hoặc tài liệu tham khảo giả. Phạm vi sử dụng được kê khai tại bảng dưới đây.", thut=0)
    bang(d, "kê khai", dict(tieu_de="Phạm vi sử dụng công cụ trí tuệ nhân tạo, Phương án B",
                           cot=["Tên công cụ", "Phạm vi hỗ trợ", "Nội dung hoặc mục có sự hỗ trợ"],
                           dong=[["", "", ""], ["", "", ""], ["", "", ""]], nguon="", rong=[3.5, 6.0, 7.5]))
    doan(d, "**Phần II. Cam kết dữ liệu gốc**", thut=0)
    doan(d, "Nhóm tác giả cam kết dữ liệu trong bài là trung thực, chính xác; toàn bộ dữ liệu là văn bản chính sách, pháp "
            "luật, hướng dẫn và công trình nghiên cứu được công bố công khai, được nhóm tác giả lưu trữ và sẵn sàng cung "
            "cấp khi Ban Biên tập yêu cầu; nhóm tác giả chịu trách nhiệm trước pháp luật và trước Hội đồng Biên tập về "
            "nội dung bản thảo.", thut=0, sau=10)
    doan(d, "Hà Nội, ngày ... tháng ... năm 2026", can=WD_ALIGN_PARAGRAPH.RIGHT, thut=0)
    doan(d, "Đại diện nhóm tác giả", can=WD_ALIGN_PARAGRAPH.RIGHT, thut=0, dam=True)
    doan(d, "(ký, ghi rõ họ tên)", can=WD_ALIGN_PARAGRAPH.RIGHT, thut=0, nghieng=True, sau=24)
    doan(d, "*" + N.TO_KHAI_GHI_CHU + "*", thut=0, co=9.5)
    d.save(RA_TO_KHAI)


def kiem_tra(d, dem):
    van_ban = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    than = van_ban.split("TÀI LIỆU THAM KHẢO")[0]
    loi = []
    for khoa, apa in zip(N.KHOA_TRICH, N.TAI_LIEU):
        if khoa not in than:
            loi.append("Tài liệu chưa được trích: " + apa[:60])
    for t in re.findall(r"\(([^()]*?\d{4}[a-z]?)\)", than):
        if re.fullmatch(r"[\d, a-z]+", t):
            continue  # trích dẫn kể: Tác giả (năm)
        for phan_ in t.split(";"):
            if not any(k.split(",")[0].split(" (")[0] in phan_ for k in N.KHOA_TRICH):
                loi.append("Trích dẫn không có trong danh mục: " + phan_)
    for t in [p.text for p in d.paragraphs]:
        if "—" in t or "–" in t:
            loi.append("Gạch dài: " + t[:60])
        if re.search(r"\b(tôi|chúng tôi)\b", t, flags=re.I):
            loi.append("Ngôi thứ nhất: " + t[:60])
    tong = sum(dem.values())
    print(f"Tiêu đề: {len(N.TIEU_DE.split())} âm tiết; tóm tắt: {dem_tu([N.TOM_TAT])}; abstract: "
          f"{dem_tu([N.ABSTRACT])} từ; từ khóa: {len(N.TU_KHOA.split(';'))} cụm; tài liệu: {len(N.TAI_LIEU)}")
    muc_tieu = {"Đặt vấn đề": 300, "Tổng quan": 800, "Phương pháp": 400, "Kết quả 4.1": 500, "Kết quả 4.2": 1400,
                "Kết quả 4.3": 400, "Bàn luận": 900, "Kết luận": 300}
    for k, v in dem.items():
        print(f"  {k:12s} {v:5d}  (khung khoảng {muc_tieu[k]})")
    bang_tu = sum(len(c.text.split()) for t in d.tables for r in t.rows for c in r.cells)
    tl = dem_tu(N.TAI_LIEU)
    print(f"  Nội dung chính {tong} từ; bảng {bang_tu}; tóm tắt song ngữ {dem_tu([N.TOM_TAT, N.ABSTRACT])}; "
          f"tài liệu {tl}; tổng {tong + bang_tu + dem_tu([N.TOM_TAT, N.ABSTRACT]) + tl}")
    for x in loi:
        print("CẢNH BÁO:", x)


if __name__ == "__main__":
    d, dem = dung_bai()
    dung_to_khai()
    print("Đã ghi:", RA_DOCX)
    print("Đã ghi:", RA_TO_KHAI)
    kiem_tra(d, dem)
