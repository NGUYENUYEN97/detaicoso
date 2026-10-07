# -*- coding: utf-8 -*-
"""Sửa Chương 1 bản cuối Báo cáo toàn văn (06/10/2026), xuất tệp có theo dõi thay đổi.

Kết hợp góp ý biên tập văn phong và các điểm chính xác đã đối chiếu:
- Luật Sở hữu trí tuệ (Văn bản hợp nhất 67/VBHN-VPQH): Điều 4, 6, 60, 86, 90, 135, 138, 141, 143, 144, 148, 198, 199;
- Luật số 93/2025/QH15: Điều 25, 27, 28, 66; Nghị định 267/2025/NĐ-CP: Điều 34;
- Luật số 125/2025/QH15: khoản 4 Điều 3, Điều 28; Nghị định 100/2026/NĐ-CP (Điều 9a Nghị định 65/2023/NĐ-CP);
- Thông tư 83/2026/TT-BGDĐT: Tiêu chí 1.1, 6.1, 6.2 và công thức quy đổi sản phẩm;
- Quyết định 1068/QĐ-TTg (bốn khâu sáng tạo, xác lập, khai thác, bảo vệ);
- Tài liệu học thuật trong thư mục: WIPO (2020), Guan (2014), Fisher (2001), Phạm Thị Thúy Hằng (2019, tr. 22),
  Tewari và Bhardwaj (2021), Bstieler và cộng sự (2015), Bradley và cộng sự (2013), Etzkowitz (2003),
  Goldfarb và Henrekson (2003), Shane (2004), Siegel và cộng sự (2007), Võ Nguyên Hoàng Phúc (2025);
- Bảng đối sánh chỉ số (Doi-sanh-chi-so-SHTT-xep-hang-kiem-dinh.xlsx) và Chiến lược Trường (UPM 3 sao năm 2020).
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from track_changes import ap_dung  # noqa: E402

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
THU_MUC = os.path.join(GOC, "Ban_cuoi", "Ban_hoan_thien_03-10-2026")
VAO = os.path.join(THU_MUC, "Bao_cao_toan_van_ban_cuoi_06-10-2026.docx")
RA = os.path.join(THU_MUC, "Bao_cao_toan_van_ban_cuoi_sua_Chuong1.docx")
MERGE = "/root/.claude/skills/synced/511a21b4-da62-4b21-9cb6-da109b6fbd84_b74c1dbc-fffb-4c9d-9a90-c57409abedcc/docx/scripts/merge_runs.py"

_D = docx.Document(VAO)
_TAT_CA = [p.text for p in _D.paragraphs]
for _t in _D.tables:
    for _r in _t.rows:
        for _c in _r.cells:
            _TAT_CA.extend(p.text for p in _c.paragraphs)


def doan(dau):
    """Trả về toàn văn đoạn duy nhất bắt đầu bằng chuỗi dau."""
    hits = sorted({t for t in _TAT_CA if t.startswith(dau)})
    if len(hits) != 1:
        raise ValueError(f"{len(hits)} đoạn bắt đầu bằng: {dau}")
    return hits[0]


def nhan_than(dau):
    """Tách đoạn dạng 'Nhãn đậm: nội dung' thành (nhãn, nội dung)."""
    t = doan(dau)
    i = t.index(":") + 1
    return t[:i], t[i:]


def thay_doan(dau, moi):
    return (doan(dau), moi)


def thay_nhan(dau, moi):
    """Bỏ nhãn đậm đầu đoạn, viết lại phần nội dung."""
    nhan, than = nhan_than(dau)
    return [(nhan, ""), (than, moi)]


# ---------------- 1.1.1 ----------------
P_MO = ("Mục này làm rõ ba khái niệm được sử dụng xuyên suốt đề tài: tài sản trí tuệ, quyền sở hữu trí tuệ và quản lý "
        "quyền sở hữu trí tuệ trong trường đại học. Mỗi khái niệm được trình bày từ định nghĩa trong tài liệu và văn bản "
        "pháp luật, sau đó xác định cách hiểu mà đề tài sử dụng.")

P_TSTT = ("Theo Tổ chức Sở hữu trí tuệ thế giới (năm 2020), sở hữu trí tuệ chỉ các sản phẩm sáng tạo của trí óc, từ tác "
          "phẩm nghệ thuật, sáng chế, chương trình máy tính đến nhãn hiệu và các dấu hiệu thương mại khác. Guan (năm 2014) "
          "lưu ý rằng chưa có một định nghĩa được thừa nhận chung về sở hữu trí tuệ. Theo cách tiếp cận vị lợi được Fisher "
          "(năm 2001) tổng kết, hệ thống sở hữu trí tuệ cần đạt điểm cân bằng tối ưu giữa quyền độc quyền nhằm khuyến khích "
          "sáng tạo và sự hạn chế mà độc quyền đó gây ra đối với việc công chúng tiếp cận tri thức. Kế thừa các cách tiếp "
          "cận trên, đề tài hiểu tài sản trí tuệ trong trường đại học là các kết quả sáng tạo trí tuệ hình thành từ hoạt "
          "động đào tạo, nghiên cứu khoa học và chuyển giao công nghệ của giảng viên, nhà khoa học, người học và người lao "
          "động của nhà trường, có thể được bảo hộ dưới một trong các đối tượng quyền sở hữu trí tuệ như sáng chế, giải "
          "pháp hữu ích, kiểu dáng công nghiệp, nhãn hiệu, chương trình máy tính, giáo trình, bài giảng số, cơ sở dữ liệu "
          "và bí mật kinh doanh.")

P_QSHTT = ("Theo khoản 1 Điều 4 Luật Sở hữu trí tuệ, quyền sở hữu trí tuệ là quyền của tổ chức, cá nhân đối với các đối "
           "tượng quyền tác giả và quyền liên quan đến quyền tác giả, quyền sở hữu công nghiệp và quyền đối với giống cây "
           "trồng. Đề tài tập trung vào hai nhóm quyền đầu, vì quyền đối với giống cây trồng ít liên quan đến các lĩnh vực "
           "đào tạo và nghiên cứu hiện nay của nhà trường. Hai nhóm quyền này khác nhau về căn cứ phát sinh và xác lập, "
           "kéo theo sự khác nhau về công cụ quản lý; nội dung này được phân tích tại mục 1.2.2.")

P_QL = ("Phạm Thị Thúy Hằng (năm 2019, tr. 22) định nghĩa quản lý hoạt động sở hữu trí tuệ trong trường đại học là "
        "\"tổng thể tác động có tổ chức, có mục đích của chủ thể quản lý trường đại học đối với hoạt động sở hữu trí tuệ "
        "nhằm tạo lập, bảo hộ, bảo vệ, khai thác và sử dụng hiệu quả sở hữu trí tuệ của trường đại học, góp phần thực "
        "hiện mục tiêu giáo dục và phát triển đất nước\". Trên cơ sở định nghĩa này, đề tài hiểu quản lý quyền sở hữu "
        "trí tuệ trong trường đại học là hoạt động có tổ chức của Hội đồng trường, Ban Giám hiệu và các đơn vị chức năng "
        "nhằm nhận diện, xác lập, khai thác và bảo vệ các quyền sở hữu trí tuệ hình thành từ hoạt động của nhà trường, "
        "bảo đảm hài hòa lợi ích của nhà trường, tác giả và các bên tham gia. Hoạt động này phải dung hòa nhiều mục tiêu: "
        "bảo vệ quyền tài sản của nhà trường và lợi ích chính đáng của tác giả, duy trì việc phổ biến tri thức qua công "
        "bố khoa học, bảo đảm liêm chính học thuật và đáp ứng yêu cầu của chuẩn chất lượng giáo dục đại học.")

# ---------------- 1.1.2 ----------------
P_DD_MO = ("Tài sản trí tuệ hình thành trong trường đại học định hướng ứng dụng có bốn đặc điểm chi phối cách thức tổ chức "
           "quản lý.")

P_DD1 = ("Một là, tính vô hình và khả năng lựa chọn hình thức bảo hộ. Nhiều kết quả có giá trị như thuật toán, bí quyết "
         "thực nghiệm hay quy trình bào chế mẫu thử tồn tại dưới dạng thông tin, chưa được ghi nhận thành văn bản và dễ bị "
         "bộc lộ trong trao đổi học thuật. Đối với các kết quả này, chủ sở hữu phải lựa chọn giữa đăng ký sáng chế và giữ "
         "bí mật; Tewari và Bhardwaj (năm 2021) cho rằng giữ bí mật chỉ phù hợp khi có thể duy trì tính bí mật trong thời "
         "gian dài và khả năng giải mã ngược công nghệ thấp. Đặc điểm này đòi hỏi nhà trường có cơ chế nhận diện sớm, "
         "khuyến khích tác giả khai báo kết quả trước khi trao đổi hoặc công bố.")

P_DD2 = ("Hai là, yêu cầu phối hợp giữa công bố khoa học và bảo hộ. Công bố kết quả là nghĩa vụ học thuật của giảng viên, "
         "trong khi sáng chế chỉ được bảo hộ nếu còn tính mới. Theo khoản 1 Điều 60 Luật Sở hữu trí tuệ, sáng chế mất tính "
         "mới nếu đã bị bộc lộ công khai trước ngày nộp đơn. Khoản 3 Điều này quy định một ngoại lệ: sáng chế không bị "
         "coi là mất tính mới nếu người có quyền đăng ký bộc lộ công khai và đơn được nộp tại Việt Nam trong thời hạn "
         "mười hai tháng kể từ ngày bộc lộ. Vì vậy, nhà trường cần sắp xếp thời điểm nộp đơn và thời điểm công bố, thay "
         "vì buộc giảng viên chọn một trong hai mục tiêu.")

P_DD3 = ("Ba là, phần lớn kết quả nghiên cứu cần được hoàn thiện thêm trước khi khai thác. Kết quả nghiên cứu trong trường "
         "đại học thường dừng ở quy mô phòng thí nghiệm hoặc mẫu thử. Để trở thành sản phẩm đáp ứng nhu cầu thị trường, "
         "kết quả cần được thử nghiệm, hoàn thiện công nghệ và sản xuất thử, nên việc định giá và lựa chọn phương thức "
         "khai thác thường có rủi ro và cần sự tham gia của doanh nghiệp.")

P_DD4 = ("Bốn là, nhiều chủ thể cùng tham gia tạo ra kết quả. Một công trình có thể có đóng góp của giảng viên, nghiên cứu "
         "viên, người học và đối tác doanh nghiệp, với các nguồn kinh phí khác nhau. Khoản 2 Điều 86 Luật Sở hữu trí tuệ "
         "quy định khi nhiều tổ chức, cá nhân cùng tạo ra hoặc cùng đầu tư để tạo ra sáng chế, kiểu dáng công nghiệp, "
         "thiết kế bố trí thì việc đăng ký chỉ được thực hiện nếu tất cả đồng ý. Do đó, quyền sở hữu, quyền đứng tên tác "
         "giả và tỷ lệ phân chia lợi ích cần được thỏa thuận bằng văn bản ngay từ khi bắt đầu nhiệm vụ.")

# ---------------- 1.1.3 ----------------
P_VT_MO = "Đối với trường đại học định hướng ứng dụng, quyền sở hữu trí tuệ có vai trò trên bốn phương diện."

P_VT1 = ("Thứ nhất, gắn đào tạo với thực tiễn. Khi giảng viên và người học cùng giải quyết các bài toán kỹ thuật, quản lý "
         "do doanh nghiệp đặt ra, kết quả tạo ra có thể trở thành tài sản trí tuệ, đồng thời trở thành học liệu và tình "
         "huống giảng dạy. Việc người học được hướng dẫn nhận diện, khai báo và tôn trọng quyền sở hữu trí tuệ cũng là một "
         "phần của năng lực nghề nghiệp.")

P_VT2 = ("Thứ hai, tạo thêm nguồn thu. Khi được khai thác qua chuyển giao quyền sử dụng, chuyển nhượng hoặc góp vốn, tài "
         "sản trí tuệ có thể tạo nguồn thu ngoài học phí. Quy mô nguồn thu này phụ thuộc vào chất lượng tài sản, nhu cầu "
         "thị trường và năng lực tổ chức khai thác của nhà trường, nên cần được xem là nguồn thu bổ sung trong dài hạn, "
         "không phải nguồn thay thế học phí.")

P_VT3 = ("Thứ ba, đáp ứng yêu cầu của chuẩn chất lượng và các bảng xếp hạng. Theo Tiêu chí 6.2 Chuẩn cơ sở giáo dục đại "
         "học ban hành kèm theo Thông tư số 83/2026/TT-BGDĐT (có hiệu lực từ ngày 15/11/2026), kết quả khoa học, công nghệ "
         "và đổi mới sáng tạo được tính theo sản phẩm quy đổi, trong đó mỗi bằng độc quyền sáng chế được tính năm sản "
         "phẩm, mỗi bằng độc quyền giải pháp hữu ích được tính ba sản phẩm, trong khi mỗi bài báo trong nước được tính một "
         "sản phẩm. Các chỉ số về sáng chế, chuyển giao và doanh nghiệp khởi nguồn cũng có mặt trong một số hệ thống kiểm "
         "định và xếp hạng, được tổng hợp tại Bảng 1.2.")

P_VT4 = ("Thứ tư, tạo động lực cho nhà khoa học. Quy chế sở hữu trí tuệ rõ ràng giúp tác giả biết trước quyền lợi của "
         "mình, gồm quyền nhân thân và mức thù lao hoặc tiền thưởng khi kết quả được khai thác. Trong hợp tác với doanh "
         "nghiệp, Bstieler và cộng sự (năm 2015) cho thấy tính linh hoạt và minh bạch của chính sách sở hữu trí tuệ của "
         "trường đại học có liên hệ với mức độ tin cậy giữa các bên, qua đó hỗ trợ quan hệ hợp tác lâu dài.")

# ---------------- 1.2 và Bảng 1.1 ----------------
P_12 = ("Đề tài phân tích nội dung quản lý quyền sở hữu trí tuệ theo bốn khâu: sáng tạo, xác lập, khai thác và bảo vệ. "
        "Cách phân chia này thống nhất với Chiến lược sở hữu trí tuệ đến năm 2030 ban hành kèm theo Quyết định số "
        "1068/QĐ-TTg, trong đó yêu cầu phát triển hệ thống sở hữu trí tuệ đồng bộ, hiệu quả ở tất cả các khâu sáng tạo, "
        "xác lập, khai thác và bảo vệ, thực thi quyền sở hữu trí tuệ. Bốn khâu không diễn ra tuần tự một chiều; Bradley "
        "và cộng sự (năm 2013) chỉ ra rằng mô hình tuyến tính chưa phản ánh đầy đủ quá trình chuyển giao công nghệ từ "
        "trường đại học, nên các khâu cần được xem là những nội dung có liên hệ qua lại. Các nhóm nội dung chính sách mà "
        "cơ sở giáo dục đại học cần quy định, cùng căn cứ pháp lý hiện hành, được tổng hợp tại Bảng 1.1.")

BANG11 = [
    ("Bảng 1.1. Khung chính sách quản lý quyền sở hữu trí tuệ",
     "Bảng 1.1. Các nhóm nội dung chính sách quản lý quyền sở hữu trí tuệ trong cơ sở giáo dục đại học và căn cứ pháp "
     "lý Việt Nam"),
    ("Trụ cột chính sách của WIPO năm 2020", "Nhóm nội dung chính sách"),
    ("1. Xác lập quyền sở hữu tài sản trí tuệ", "1. Quyền sở hữu và quyền đăng ký đối với kết quả sáng tạo"),
    ("- Điều 86 Luật Sở hữu trí tuệ<br>",
     "- Điểm b, điểm c khoản 1 Điều 86 Luật Sở hữu trí tuệ\n- Khoản 2 Điều 25 Luật số 93/2025/QH15"),
    ("Cơ sở giáo dục đại học là chủ sở hữu quyền tài sản đối với sáng chế",
     "Tổ chức đầu tư kinh phí, phương tiện vật chất cho tác giả dưới hình thức giao việc, thuê việc có quyền đăng ký "
     "sáng chế, kiểu dáng công nghiệp, thiết kế bố trí, trừ trường hợp các bên có thỏa thuận khác. Tổ chức chủ trì nhiệm "
     "vụ sử dụng ngân sách nhà nước được tự động giao quyền quản lý, sử dụng, quyền sở hữu phần kết quả tương ứng với "
     "kinh phí ngân sách, không phải bồi hoàn, trừ các trường hợp luật định."),
    ("Minh định quyền sở hữu tài sản trí tuệ thuộc về nhà trường",
     "Quy định rõ quyền sở hữu, quyền đăng ký đối với kết quả tạo ra từ nhiệm vụ được giao, từ kinh phí của nhà trường "
     "và từ hợp tác với bên ngoài, trong quy chế nội bộ và hợp đồng nghiên cứu."),
    ("- Điều 60 và Điều 90 Luật Sở hữu trí tuệ<br>",
     "- Điều 60 và Điều 90 Luật Sở hữu trí tuệ\n- Điều 9a Nghị định số 65/2023/NĐ-CP, được bổ sung bởi Nghị định số "
     "100/2026/NĐ-CP"),
    ("Sáng chế và giải pháp hữu ích bắt buộc phải đáp ứng tính mới",
     "Sáng chế phải có tính mới tại ngày nộp đơn, trừ trường hợp người có quyền đăng ký tự bộc lộ và nộp đơn tại Việt "
     "Nam trong thời hạn mười hai tháng; đơn nộp trước được xem xét trước. Chủ sở hữu lập, lưu giữ và cập nhật hằng năm "
     "Danh mục quyền sở hữu trí tuệ phục vụ quản trị nội bộ."),
    ("Ban hành mẫu phiếu khai báo sáng kiến nội bộ",
     "Ban hành mẫu phiếu khai báo kết quả sáng tạo; quy định trình tự xem xét khả năng bảo hộ và thời điểm nộp đơn "
     "trước khi công bố; lập và cập nhật Danh mục quyền sở hữu trí tuệ của nhà trường."),
    ("3. Thẩm định bảo hộ và chi trả chi phí", "3. Kinh phí đăng ký, bảo hộ và quản lý quyền"),
    ("- Luật Sở hữu trí tuệ<br>",
     "- Điểm b khoản 2 Điều 66 Luật số 93/2025/QH15\n- Điểm d khoản 3 Điều 28 Luật số 125/2025/QH15"),
    ("Xác lập quyền sở hữu công nghiệp theo nguyên tắc nộp đơn đầu tiên",
     "Doanh nghiệp, tổ chức, đơn vị sự nghiệp được trích lập quỹ phát triển khoa học và công nghệ và được dùng quỹ để "
     "hỗ trợ nghiên cứu, đổi mới sáng tạo và đăng ký, bảo hộ, quản lý, khai thác quyền sở hữu trí tuệ. Cơ sở giáo dục "
     "đại học có trách nhiệm thành lập và vận hành quỹ này theo quy định."),
    ("Bố trí danh mục kinh phí thường niên chuyên biệt",
     "Bố trí kinh phí hằng năm cho tra cứu, nộp đơn, duy trì hiệu lực văn bằng bảo hộ và thuê tổ chức dịch vụ đại diện "
     "sở hữu công nghiệp khi cần."),
    ("- Chương X Luật Sở hữu trí tuệ<br>",
     "- Chương X Luật Sở hữu trí tuệ\n- Điều 27 Luật số 93/2025/QH15\n- Khoản 1, điểm d khoản 2 Điều 28 Luật số "
     "125/2025/QH15"),
    ("Cho phép trường định giá, góp vốn bằng tài sản trí tuệ",
     "Chủ sở hữu tự quyết định việc thương mại hóa; tổ chức được giao quyền tự quyết định hình thức, giá, phân chia lợi "
     "nhuận và phương án góp vốn. Cơ sở giáo dục đại học được thành lập doanh nghiệp khoa học và công nghệ, doanh "
     "nghiệp quản lý tài sản trí tuệ và được định giá, góp vốn, phân chia lợi ích từ tài sản trí tuệ."),
    ("Chuẩn hóa quy trình định giá tài sản trí tuệ nội bộ",
     "Ban hành quy trình định giá tài sản trí tuệ, mẫu hợp đồng chuyển quyền sử dụng, chuyển nhượng và cơ chế tham gia "
     "doanh nghiệp khởi nguồn từ kết quả nghiên cứu của nhà trường."),
    ("5. Phân chia lợi ích tài chính và đòn bẩy động lực", "5. Phân chia lợi ích và chính sách khuyến khích"),
    ("- Điều 135 Luật Sở hữu trí tuệ<br>",
     "- Điều 135 Luật Sở hữu trí tuệ\n- Khoản 2, điểm a khoản 3 Điều 28 Luật số 93/2025/QH15"),
    ("Quy định mức thù lao tối thiểu cho tác giả",
     "Chủ sở hữu trả thù lao cho tác giả sáng chế, kiểu dáng công nghiệp, thiết kế bố trí theo thỏa thuận; nếu không "
     "có thỏa thuận, mức thù lao là 10% lợi nhuận trước thuế tương ứng khi tự sử dụng hoặc 15% số tiền nhận được mỗi lần "
     "thanh toán khi chuyển giao quyền sử dụng. Với phần kết quả sử dụng ngân sách nhà nước, tác giả được thưởng tối "
     "thiểu 30% lợi nhuận thu được; với phần không sử dụng ngân sách nhà nước, chủ sở hữu tự quyết định."),
    ("Cụ thể hóa tỷ lệ chia thưởng tài chính trong quy chế nội bộ",
     "Quy định tỷ lệ phân chia lợi ích giữa tác giả, đơn vị và nhà trường trong quy chế nội bộ, phân biệt theo nguồn "
     "kinh phí; kết hợp tiền thưởng với quy đổi giờ nghiên cứu khoa học."),
    ("6. Quản lý xung đột lợi ích và liêm chính học thuật", "6. Bảo mật, liêm chính học thuật và xử lý vi phạm"),
    ("- Điểm c Khoản 4 Điều 3 và Điều 28 Luật số 125/2025/QH15",
     "- Khoản 4 Điều 3 và điểm a khoản 3 Điều 28 Luật số 125/2025/QH15\n- Điều 127 Luật Sở hữu trí tuệ\n- Tiêu chí 1.1 "
     "Chuẩn cơ sở giáo dục đại học ban hành kèm theo Thông tư số 83/2026/TT-BGDĐT"),
    ("Cơ sở giáo dục đại học bắt buộc ban hành quy chế liêm chính học thuật",
     "Liêm chính học thuật bao gồm yêu cầu không sao chép, làm giả dữ liệu, vi phạm quyền tác giả, quyền sở hữu trí "
     "tuệ; cơ sở giáo dục đại học phải tuân thủ quy tắc đạo đức, liêm chính trong nghiên cứu. Quy định về sở hữu trí "
     "tuệ, liêm chính khoa học, liêm chính học thuật thuộc nhóm văn bản nội bộ bắt buộc. Hành vi tiếp cận, bộc lộ trái "
     "phép bí mật kinh doanh là hành vi xâm phạm."),
    ("Thành lập hội đồng liêm chính học thuật, ký cam kết bảo mật",
     "Ban hành quy định về sở hữu trí tuệ và liêm chính học thuật; yêu cầu cam kết bảo mật đối với người tham gia "
     "nhiệm vụ; quy định trình tự xử lý tranh chấp, vi phạm."),
    ("(Nguồn: Nhóm nghiên cứu tổng hợp từ Bộ công cụ chính sách",
     "(Nguồn: Nhóm nghiên cứu tổng hợp từ hệ thống văn bản quy phạm pháp luật Việt Nam giai đoạn 2025 - 2026)"),
]

# ---------------- 1.2.1 ----------------
P_ST1 = ("Khâu sáng tạo tạo ra nguồn tài sản trí tuệ cho các khâu sau. Nội dung quản lý ở khâu này gồm định hướng nghiên "
         "cứu, bố trí nguồn lực và chính sách khuyến khích, nhằm tăng số kết quả có khả năng bảo hộ và ứng dụng. Etzkowitz "
         "(năm 2003) chỉ ra rằng các nhóm nghiên cứu trong trường đại học mang nhiều đặc điểm của doanh nghiệp, nhất là "
         "khi kinh phí nghiên cứu được phân bổ theo cơ chế cạnh tranh; vì vậy, nhóm nghiên cứu là đơn vị cơ sở cần được "
         "hỗ trợ ở khâu sáng tạo.")

P_ST2 = ("Trên cơ sở đó, quản lý khâu sáng tạo trong trường đại học định hướng ứng dụng gồm ba nội dung. Một là, định "
         "hướng đề tài theo nhu cầu của doanh nghiệp và địa phương, ưu tiên các bài toán có khả năng tạo ra giải pháp kỹ "
         "thuật. Hai là, đa dạng hóa nguồn kinh phí nghiên cứu, kết hợp ngân sách nhà nước, kinh phí của nhà trường và "
         "kinh phí từ doanh nghiệp. Ba là, xây dựng chính sách khuyến khích cân đối giữa quy đổi giờ nghiên cứu khoa học, "
         "tiền thưởng và thù lao khi kết quả được khai thác. Goldfarb và Henrekson (năm 2003) so sánh Thụy Điển và Hoa Kỳ, "
         "cho rằng chính sách áp đặt từ trên xuống cùng môi trường học thuật không khuyến khích giảng viên tham gia là một "
         "nguyên nhân khiến kết quả thương mại hóa ở Thụy Điển thấp, trong khi cạnh tranh giữa các trường ở Hoa Kỳ tạo điều "
         "kiện để giảng viên tương tác với doanh nghiệp.")

# ---------------- 1.2.2 ----------------
P_XL1 = ("Xác lập quyền là khâu chuyển kết quả sáng tạo thành đối tượng được pháp luật bảo hộ. Theo Điều 6 Luật Sở hữu "
         "trí tuệ, quyền tác giả phát sinh kể từ khi tác phẩm được sáng tạo và được thể hiện dưới một hình thức vật chất "
         "nhất định, không phụ thuộc vào việc đăng ký; quyền sở hữu công nghiệp đối với sáng chế, kiểu dáng công nghiệp, "
         "thiết kế bố trí, nhãn hiệu được xác lập trên cơ sở quyết định cấp văn bằng bảo hộ; quyền đối với bí mật kinh "
         "doanh được xác lập trên cơ sở có được một cách hợp pháp và thực hiện việc bảo mật. Vì vậy, nội dung quản lý ở "
         "khâu này gồm xác định chủ thể có quyền, lựa chọn hình thức bảo hộ, lưu giữ chứng cứ sáng tạo đối với tác phẩm, "
         "nộp đơn đúng thời điểm đối với sáng chế, kiểu dáng và áp dụng biện pháp bảo mật đối với bí mật kinh doanh.")

P_XL2 = ("Về chủ thể có quyền đăng ký, điểm b khoản 1 Điều 86 Luật Sở hữu trí tuệ quy định tổ chức đầu tư kinh phí, "
         "phương tiện vật chất cho tác giả dưới hình thức giao việc, thuê việc có quyền đăng ký sáng chế, kiểu dáng công "
         "nghiệp, thiết kế bố trí, trừ trường hợp các bên có thỏa thuận khác. Như vậy, khi giảng viên tạo ra kết quả trong "
         "khuôn khổ nhiệm vụ được giao và bằng kinh phí, cơ sở vật chất của nhà trường, quyền đăng ký thuộc về nhà trường, "
         "còn giảng viên là tác giả. Đối với nhiệm vụ sử dụng ngân sách nhà nước, khoản 2 Điều 25 Luật Khoa học, công nghệ "
         "và đổi mới sáng tạo số 93/2025/QH15 quy định tổ chức chủ trì được Nhà nước tự động giao quyền quản lý, sử dụng, "
         "quyền sở hữu phần kết quả tương ứng với kinh phí ngân sách, không phải làm thủ tục giao quyền và không phải bồi "
         "hoàn, trừ một số trường hợp như nhiệm vụ quốc phòng, an ninh hoặc nhiệm vụ do Nhà nước đặt hàng có yêu cầu nắm "
         "giữ quyền. Tương ứng, điểm c khoản 1 Điều 86 Luật Sở hữu trí tuệ, được bổ sung bởi Luật số 131/2025/QH15 và có "
         "hiệu lực từ ngày 01 tháng 4 năm 2026, ghi nhận tổ chức được giao quyền này có quyền đăng ký sáng chế, kiểu dáng "
         "công nghiệp, thiết kế bố trí là kết quả của nhiệm vụ. Quy định mới giúp nhà trường chủ động hơn trong việc đăng "
         "ký và thương mại hóa kết quả nhiệm vụ sử dụng ngân sách nhà nước. Kinh nghiệm của Hoa Kỳ cho thấy tác động của "
         "việc trao quyền cho trường đại học: Shane (năm 2004) chỉ ra rằng sau khi Đạo luật Bayh-Dole có hiệu lực, tỷ "
         "trọng sáng chế của các trường đại học tăng ở những lĩnh vực mà cấp phép là cơ chế chuyển giao hiệu quả.")

P_XL3 = ("Về quản trị nội bộ, nhà trường cần quy định trình tự khai báo kết quả, xem xét khả năng bảo hộ và quyết định "
         "thời điểm nộp đơn trước khi tác giả công bố hoặc trao đổi kết quả với bên ngoài. Bên cạnh đó, Điều 9a Nghị định "
         "số 65/2023/NĐ-CP, được bổ sung bởi Nghị định số 100/2026/NĐ-CP (có hiệu lực từ ngày 01 tháng 4 năm 2026), quy "
         "định chủ sở hữu quyền sở hữu trí tuệ lập, lưu giữ Danh mục quyền sở hữu trí tuệ chưa đủ điều kiện ghi nhận giá "
         "trị tài sản trong sổ kế toán để phục vụ quản trị nội bộ. Danh mục gồm thông tin về đối tượng, tình trạng pháp "
         "lý, thời hạn bảo hộ, các mốc nộp phí, tác giả, nguồn gốc hình thành, chi phí, tình trạng khai thác và giá trị "
         "ước tính (nếu có), được rà soát, cập nhật hằng năm hoặc khi có thay đổi. Danh mục này là công cụ để nhà trường "
         "theo dõi tài sản trí tuệ từ khi hình thành đến khi khai thác.")

# ---------------- 1.2.3 ----------------
P_KT0 = "Khai thác là khâu hiện thực hóa giá trị của tài sản trí tuệ. Các hình thức khai thác chủ yếu gồm:"

P_KT2 = ("Thứ hai, chuyển quyền sử dụng: theo Điều 141 và Điều 143 Luật Sở hữu trí tuệ, nhà trường cho phép tổ chức, cá "
         "nhân khác sử dụng đối tượng sở hữu công nghiệp trong phạm vi, thời hạn thỏa thuận, theo hình thức độc quyền "
         "hoặc không độc quyền, và nhận khoản thanh toán theo thỏa thuận. Hợp đồng sử dụng không được có các điều khoản "
         "hạn chế bất hợp lý quyền của bên được chuyển quyền (khoản 2 Điều 144) và, trừ hợp đồng sử dụng nhãn hiệu, phải "
         "được đăng ký mới có giá trị pháp lý đối với bên thứ ba (khoản 3 Điều 148).")

P_KT3 = ("Thứ ba, chuyển nhượng quyền sở hữu công nghiệp: theo Điều 138 Luật Sở hữu trí tuệ, nhà trường chuyển giao quyền "
         "sở hữu đối tượng sở hữu công nghiệp cho bên nhận và nhận khoản thanh toán theo thỏa thuận. Đối với các đối tượng "
         "được xác lập trên cơ sở đăng ký, hợp đồng chuyển nhượng chỉ có hiệu lực khi đã được đăng ký tại cơ quan quản lý "
         "nhà nước về quyền sở hữu công nghiệp (khoản 1 Điều 148).")

P_KT4 = ("Thứ tư, thành lập hoặc góp vốn vào doanh nghiệp: khoản 1 Điều 28 Luật Giáo dục đại học số 125/2025/QH15 cho "
         "phép cơ sở giáo dục đại học thành lập doanh nghiệp khoa học và công nghệ, doanh nghiệp quản lý tài sản trí tuệ "
         "và đầu tư vốn vào doanh nghiệp khoa học và công nghệ; khoản 2 Điều 27 Luật số 93/2025/QH15 cho phép tổ chức "
         "được giao quyền tự quyết định phương án góp vốn bằng kết quả nghiên cứu, xác định giá góp vốn và tỷ lệ vốn góp. "
         "Siegel và cộng sự (năm 2007) nhấn mạnh trường đại học cần xây dựng chiến lược chuyển giao công nghệ nhất quán "
         "và khả thi.")

P_LI1 = ("Về phân chia lợi ích, cần phân biệt ba căn cứ. Thứ nhất, Điều 135 Luật Sở hữu trí tuệ quy định chủ sở hữu sáng "
         "chế, kiểu dáng công nghiệp, thiết kế bố trí trả thù lao cho tác giả theo thỏa thuận; chỉ khi không có thỏa "
         "thuận, mức thù lao mới là 10% lợi nhuận trước thuế tương ứng với giá trị đóng góp của đối tượng nếu chủ sở hữu "
         "tự sử dụng, hoặc 15% tổng số tiền nhận được trong mỗi lần thanh toán nếu chuyển giao quyền sử dụng. Đây là mức "
         "áp dụng khi các bên không thỏa thuận, không phải mức tối thiểu bắt buộc.")

P_LI2 = ("Thứ hai, Điều 28 Luật số 93/2025/QH15 phân biệt theo nguồn kinh phí. Với phần lợi nhuận tương ứng với kết quả "
         "không sử dụng ngân sách nhà nước, chủ sở hữu tự quyết định việc xử lý lợi nhuận, kể cả mức thưởng cho tác giả "
         "(khoản 2). Với phần tương ứng với kinh phí ngân sách nhà nước, tổ chức chủ trì dùng lợi nhuận sau thuế để thưởng "
         "cho tác giả tối thiểu 30% lợi nhuận thu được, thưởng cho người trực tiếp tổ chức thương mại hóa và tái đầu tư "
         "cho khoa học, công nghệ, đổi mới sáng tạo (khoản 3). Tổ chức trung gian, môi giới được hưởng một phần lợi nhuận; "
         "theo khoản 2 Điều 34 Nghị định số 267/2025/NĐ-CP, mức này tối thiểu là 10% lợi nhuận, trừ trường hợp các bên có "
         "thỏa thuận khác.")

P_LI3 = ("Thứ ba, đối với kết quả do nhà trường tự đầu tư, quy chế nội bộ là căn cứ chủ yếu để xác định tỷ lệ chia cho "
         "tác giả, đơn vị và nhà trường. Điểm d khoản 2 Điều 28 Luật Giáo dục đại học số 125/2025/QH15 xác định việc định "
         "giá, xác lập quyền sở hữu, khai thác, góp vốn, phân chia lợi ích từ tài sản trí tuệ là một hình thức phát triển "
         "tiềm lực khoa học, công nghệ của cơ sở giáo dục đại học. Quy chế cũng cần xác định rõ quyền của tác giả đối với "
         "tác phẩm mà nhà trường nắm giữ quyền tài sản, một vấn đề còn ít được đề cập trong quy chế sở hữu trí tuệ của các "
         "trường đại học Việt Nam (Võ Nguyên Hoàng Phúc, năm 2025). Phần nguồn thu nhà trường giữ lại có thể được đưa vào "
         "Quỹ phát triển khoa học và công nghệ; theo điểm b khoản 2 Điều 66 Luật số 93/2025/QH15, quỹ được dùng để hỗ trợ "
         "đăng ký, bảo hộ, quản lý, khai thác quyền sở hữu trí tuệ, tạo nguồn tái đầu tư cho các chu kỳ nghiên cứu tiếp "
         "theo.")

# ---------------- 1.2.4 ----------------
P_BV0 = ("Bảo vệ quyền là khâu duy trì hiệu lực và giá trị của tài sản trí tuệ đã được xác lập. Nội dung quản lý ở khâu "
         "này gồm bảo mật thông tin chưa công bố, theo dõi thời hạn và nộp phí duy trì hiệu lực văn bằng bảo hộ, lưu giữ "
         "chứng cứ về quá trình sáng tạo, phát hiện và xử lý hành vi xâm phạm, phối hợp với cơ quan có thẩm quyền khi cần.")

P_BV1 = ("Khi phát sinh hành vi xâm phạm, theo Điều 198 Luật Sở hữu trí tuệ, chủ thể quyền được tự bảo vệ bằng cách áp "
         "dụng biện pháp công nghệ; yêu cầu người vi phạm chấm dứt hành vi, xin lỗi, cải chính công khai, bồi thường thiệt "
         "hại; yêu cầu cơ quan nhà nước có thẩm quyền xử lý; hoặc khởi kiện ra tòa án, trọng tài. Theo Điều 199, tùy tính "
         "chất, mức độ, hành vi xâm phạm có thể bị xử lý bằng biện pháp dân sự, hành chính hoặc hình sự; biện pháp dân sự "
         "được quy định tại Điều 202, hành vi bị xử phạt hành chính tại Điều 211 Luật Sở hữu trí tuệ, còn trách nhiệm hình "
         "sự được xác định theo Điều 225, Điều 226 Bộ luật Hình sự.")

P_BV2 = ("Cần phân biệt hành vi xâm phạm quyền sở hữu trí tuệ với vi phạm liêm chính học thuật. Theo khoản 4 Điều 3 Luật "
         "Giáo dục đại học số 125/2025/QH15, liêm chính học thuật bao gồm yêu cầu không vi phạm quyền tác giả, quyền sở "
         "hữu trí tuệ, nhưng phạm vi rộng hơn, bao trùm cả việc sao chép, xuyên tạc, làm giả dữ liệu. Một hành vi đạo văn "
         "có thể đồng thời là xâm phạm quyền tác giả và vi phạm liêm chính, nên cần được xử lý theo cả quy định pháp luật "
         "và quy định nội bộ. Tiêu chí 1.1 Chuẩn cơ sở giáo dục đại học ban hành kèm theo Thông tư số 83/2026/TT-BGDĐT (có "
         "hiệu lực từ ngày 15/11/2026) xếp quy định về sở hữu trí tuệ, liêm chính khoa học, liêm chính học thuật vào nhóm "
         "văn bản nội bộ bắt buộc. Dữ liệu phục vụ đánh giá mức độ đáp ứng chuẩn được cập nhật lên HEMIS, và cơ sở đào tạo "
         "công bố kết quả tự đánh giá trước ngày 31 tháng 5 hằng năm.")

P_B12 = ("Để đánh giá kết quả quản lý quyền sở hữu trí tuệ, đề tài đối chiếu các chỉ số về sở hữu trí tuệ và chuyển giao "
         "công nghệ trong một số văn bản quy phạm pháp luật, chuẩn kiểm định và bảng xếp hạng đại học. Bảng 1.2 tổng hợp "
         "các chỉ số này, làm căn cứ tham chiếu cho phần đánh giá thực trạng ở Chương 2. Mức độ ràng buộc của các hệ thống "
         "khác nhau: chuẩn cơ sở giáo dục đại học và tiêu chuẩn kiểm định là yêu cầu pháp lý, còn các bảng xếp hạng chỉ có "
         "giá trị tham khảo.")

BANG12 = [
    ("Bảng 1.2. Bảng đối sánh các chỉ số",
     "Bảng 1.2. Các chỉ số sở hữu trí tuệ và chuyển giao công nghệ trong chuẩn chất lượng, kiểm định và xếp hạng đại "
     "học"),
    ("Ý nghĩa quản trị đối với cơ sở giáo dục đại học định hướng ứng dụng",
     "Ý nghĩa đối với cơ sở giáo dục đại học định hướng ứng dụng"),
    ("Tiêu chí 4.2: Số lượng sáng chế",
     "Tiêu chí 6.1: tỷ trọng thu từ hoạt động khoa học, công nghệ và đổi mới sáng tạo trên tổng thu.\nTiêu chí 6.2: số "
     "sản phẩm khoa học, công nghệ quy đổi bình quân trên một giảng viên quy đổi, trong đó bằng độc quyền sáng chế tính "
     "5, bằng độc quyền giải pháp hữu ích tính 3.\nTiêu chí 1.1: có quy định về sở hữu trí tuệ, liêm chính khoa học, "
     "liêm chính học thuật."),
    ("Đạt tối thiểu 0,02 văn bằng",
     "Tiêu chí 6.1: không thấp hơn 5%, trung bình 3 năm, áp dụng với cơ sở có đào tạo tiến sĩ.\nTiêu chí 6.2: không "
     "thấp hơn 0,3 sản phẩm quy đổi/giảng viên/năm (0,6 đối với cơ sở có đào tạo tiến sĩ)."),
    ("Điều kiện tiên quyết khẳng định tư cách pháp lý",
     "Là yêu cầu bắt buộc; sáng chế và giải pháp hữu ích có hệ số quy đổi cao nên tác động trực tiếp đến khả năng đạt "
     "Tiêu chí 6.2."),
    ("Tiêu chí 7.2: Hiệu quả nghiên cứu",
     "Tiêu chí 7.2: chỉ tiêu nghiên cứu có sáng chế, bản quyền.\nTiêu chí 7.4: hệ thống dữ liệu về tài sản trí tuệ.\n"
     "Tiêu chí 13.2: loại hình và số lượng tài sản trí tuệ.\nTiêu chí 13.3: khởi nghiệp, ươm tạo và thương mại hóa."),
    ("Chiếm 6,7% tổng số tiêu chí kiểm định",
     "4 trên 60 tiêu chí (khoảng 6,7%), thuộc 2 trên 15 tiêu chuẩn."),
    ("Cung cấp bằng chứng thực chất",
     "Đòi hỏi nhà trường có chỉ tiêu và cơ sở dữ liệu theo dõi tài sản trí tuệ tập trung để làm minh chứng."),
    ("Việt Nam và Đông Nam Á (Thành Đô đạt 4 sao)",
     "Việt Nam và Đông Nam Á (Trường Đại học Thành Đô đạt 3 sao theo định hướng ứng dụng năm 2020)"),
    ("Tiêu chuẩn Đổi mới sáng tạo (4 tiêu chí",
     "Tiêu chuẩn Đổi mới sáng tạo (4 tiêu chí); Tiêu chuẩn Hệ sinh thái đổi mới sáng tạo (4 tiêu chí)."),
    ("Chiếm 17,0% tổng số điểm",
     "170 trên 1.000 điểm (17%), gồm Tiêu chuẩn Đổi mới sáng tạo khoảng 11% và Tiêu chuẩn Hệ sinh thái đổi mới sáng "
     "tạo khoảng 6%."),
    ("Định vị trực tiếp uy tín học thuật",
     "Là căn cứ tham chiếu khi nhà trường đăng ký đánh giá, nâng hạng sao."),
    ("Bảng xếp hạng đại học thế giới THE Impact Rankings",
     "THE Impact Rankings - Mục tiêu phát triển bền vững số 9"),
    ("Số bằng sáng chế trích dẫn nghiên cứu của trường;",
     "Số bằng sáng chế trích dẫn nghiên cứu của trường (15,4%); thu nhập nghiên cứu từ doanh nghiệp (38,4%); số doanh "
     "nghiệp khởi nguồn từ trường (34,6%)."),
    ("Chiếm 88,4% tổng số điểm của Mục tiêu SDG 9",
     "88,4% điểm của SDG 9. Nếu SDG 9 thuộc ba mục tiêu có điểm cao nhất của trường (mỗi mục tiêu chiếm 26% điểm "
     "tổng), phần này tương đương khoảng 23% điểm tổng."),
    ("Thước đo hội nhập quốc tế cao nhất",
     "Phản ánh mức đóng góp của tài sản trí tuệ vào phát triển công nghiệp; chỉ có ý nghĩa khi trường tham gia và lựa "
     "chọn SDG 9."),
    ("Bảng xếp hạng viện nghiên cứu và đại học SCImago", "SCImago Institutions Rankings"),
    ("Toàn cầu, trích xuất dữ liệu Scopus", "Toàn cầu, dữ liệu Scopus và PATSTAT"),
    ("Trụ cột Đổi mới sáng tạo: Tỷ lệ số bằng sáng chế",
     "Nhóm Đổi mới sáng tạo: số đơn sáng chế (10%); số công bố được trích dẫn trong sáng chế (10%); tác động công "
     "nghệ của công bố (10%)."),
    ("Chiếm 30,0% tổng trọng số", "30% tổng trọng số."),
    ("Khẳng định nghiên cứu khoa học của nhà trường",
     "Cho thấy mức độ kết quả nghiên cứu của trường được sử dụng trong phát triển công nghệ."),
    ("Quy định hoạt động KHCN trong trường đại học",
     "Quy định về hoạt động khoa học và công nghệ trong cơ sở giáo dục đại học (Nghị định số 109/2022/NĐ-CP)"),
    ("Toàn quốc (Khung quy chuẩn)", "Toàn quốc (tiêu chí nhóm nghiên cứu mạnh)"),
    ("Đạt tối thiểu 01 sáng chế/năm",
     "Một trong các tiêu chí: bình quân ít nhất 01 bằng sáng chế hoặc 02 bằng giải pháp hữu ích mỗi năm; hoặc trong 5 "
     "năm chuyển giao ít nhất 05 công nghệ hoặc thương mại hóa ít nhất 05 sản phẩm, hoặc có 01 sản phẩm quốc gia."),
    ("Mốc chuẩn định lượng then chốt để thành lập",
     "Là căn cứ xác định nhóm nghiên cứu mạnh và định hướng đầu tư cho nhóm nghiên cứu."),
    ("(Nguồn: Nhóm nghiên cứu tổng hợp và đối sánh",
     "(Nguồn: Nhóm nghiên cứu tổng hợp từ Thông tư số 83/2026/TT-BGDĐT, Thông tư số 20/2026/TT-BGDĐT, Nghị định số "
     "109/2022/NĐ-CP và phương pháp đánh giá của UPM, THE Impact Rankings, SCImago Institutions Rankings. Các chỉ số "
     "của bảng xếp hạng chỉ có giá trị tham khảo, không thay thế yêu cầu bắt buộc của Chuẩn cơ sở giáo dục đại học)"),
]

P_TK = ("Chương 1 đã làm rõ các khái niệm tài sản trí tuệ, quyền sở hữu trí tuệ và quản lý quyền sở hữu trí tuệ trong "
        "trường đại học; chỉ ra bốn đặc điểm của tài sản trí tuệ trong trường đại học định hướng ứng dụng và bốn vai trò "
        "của quyền sở hữu trí tuệ đối với loại hình trường này. Nội dung quản lý được phân tích theo bốn khâu sáng tạo, "
        "xác lập, khai thác và bảo vệ, gắn với các quy định mới của Luật số 93/2025/QH15, Luật số 131/2025/QH15, Luật số "
        "125/2025/QH15, Nghị định số 100/2026/NĐ-CP và Thông tư số 83/2026/TT-BGDĐT. Bảng 1.1 tổng hợp sáu nhóm nội dung "
        "chính sách mà cơ sở giáo dục đại học cần quy định; Bảng 1.2 tổng hợp các chỉ số sở hữu trí tuệ trong chuẩn chất "
        "lượng, kiểm định và xếp hạng đại học. Các nội dung này là khung để đánh giá thực trạng quản lý quyền sở hữu trí "
        "tuệ tại Trường Đại học Thành Đô ở Chương 2.")


def danh_sach_sua():
    sua = [
        thay_doan("Trong nền kinh tế tri thức và bối cảnh tự chủ đại học", P_MO),
        *thay_nhan("Tài sản trí tuệ trong trường đại học:", P_TSTT),
        *thay_nhan("Quyền sở hữu trí tuệ trong môi trường đại học:", P_QSHTT),
        *thay_nhan("Quản lý quyền sở hữu trí tuệ trong trường đại học:", P_QL),
        thay_doan("Tài sản trí tuệ hình thành trong môi trường giáo dục đại học", P_DD_MO),
        thay_doan("Một là, tính vô hình và sự tồn tại", P_DD1),
        thay_doan("Hai là, tính lưỡng dụng và xung đột", P_DD2),
        thay_doan("Ba là, tính phụ thuộc vào chu trình", P_DD3),
        thay_doan("Bốn là, tính đa dạng và phức tạp", P_DD4),
        thay_doan("Đối với các trường đại học định hướng ứng dụng, quyền sở hữu trí tuệ giữ vai trò", P_VT_MO),
        thay_doan("Thứ nhất, nâng cao chất lượng đào tạo", P_VT1),
        thay_doan("Thứ hai, đa dạng hóa nguồn thu", P_VT2),
        thay_doan("Thứ ba, khẳng định vị thế", P_VT3),
        thay_doan("Thứ tư, thiết lập hành lang pháp lý", P_VT4),
        thay_doan("Quản lý quyền sở hữu trí tuệ trong cơ sở giáo dục đại học bao quát", P_12),
    ]
    sua += [thay_doan(dau, moi) for dau, moi in BANG11]
    sua += [
        thay_doan("Dưới đây là nội dung quản lý chuyên sâu", "Nội dung quản lý ở từng khâu được trình bày dưới đây."),
        ("1.2.1. Sáng tạo quyền sở hữu trí tuệ", "1.2.1. Sáng tạo tài sản trí tuệ", "tat_ca"),
        thay_doan("Khâu sáng tạo là điểm khởi đầu", P_ST1),
        thay_doan("Để nâng cao hiệu quả khâu sáng tạo", P_ST2),
        thay_doan("Xác lập quyền sở hữu trí tuệ là khâu pháp lý", P_XL1),
        thay_doan("Về quyền nộp đơn đăng ký bảo hộ", P_XL2),
        thay_doan("Về quy trình quản trị nội bộ", P_XL3),
        thay_doan("Khai thác quyền sở hữu trí tuệ là khâu hiện thực hóa", P_KT0),
        thay_doan("Thứ hai, chuyển quyền sử dụng", P_KT2),
        thay_doan("Thứ ba, chuyển nhượng quyền sở hữu công nghiệp", P_KT3),
        thay_doan("Thứ tư, thành lập doanh nghiệp khởi nguồn", P_KT4),
        # chèn hai đoạn mới trước khi thay đoạn neo (đoạn đã thay không còn chữ để neo)
        ("Về cơ chế phân chia lợi ích tài chính", [(None, P_LI2), (None, P_LI3)], "cac_doan_sau"),
        thay_doan("Về cơ chế phân chia lợi ích tài chính", P_LI1),
        thay_doan("Bảo vệ quyền sở hữu trí tuệ là khâu bảo đảm", P_BV0),
        thay_doan("Khi phát sinh hành vi xâm phạm quyền", P_BV1),
        thay_doan("Đồng thời, theo Thông tư số 83/2026/TT-BGDĐT", P_BV2),
        thay_doan("Để lượng hóa hiệu quả quản lý", P_B12),
    ]
    sua += [thay_doan(dau, moi) for dau, moi in BANG12]
    sua.append(thay_doan("Chương 1 đã hệ thống hóa toàn diện", P_TK))
    return sua


if __name__ == "__main__":
    SUA = danh_sach_sua()
    n = ap_dung(VAO, RA, SUA, author="Claude", merge_runs=MERGE if os.path.exists(MERGE) else None)
    print(f"Đã áp dụng {n} chỉnh sửa -> {RA}")
