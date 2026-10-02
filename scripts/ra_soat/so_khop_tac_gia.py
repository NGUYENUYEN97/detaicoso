# -*- coding: utf-8 -*-
"""Tái lập phép so khớp tác giả bài báo với danh sách nhân sự năm 2026 (chỉ xuất số liệu tổng hợp).

    python3 scripts/ra_soat/so_khop_tac_gia.py

Ba quy tắc chuẩn hóa tên, áp dụng thống nhất cho tác giả và nhân sự (chữ thường; bỏ dấu chấm, gạch nối; bỏ học
hàm, học vị đứng trước tên):
- A, quy tắc chính: giữ dấu tiếng Việt, so khớp đúng thứ tự âm tiết;
- B, độ nhạy: bỏ dấu, đ -> d, so khớp đúng thứ tự;
- C, độ nhạy: bỏ dấu, so khớp theo tập hợp âm tiết, nhận cả cách viết họ sau tên.
Nhân sự trùng họ tên sau chuẩn hóa được tính là một tên, vì không thể phân biệt từ danh mục bài báo.
Đơn vị đếm người là "tên chuẩn hóa"; đơn vị đếm bài là dòng bài báo trong danh mục.
Không ghi ra bất kỳ thông tin cá nhân nào.
"""
import collections
import os
import re
import unicodedata

import openpyxl
import pymupdf

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TD = os.path.join(GOC, "Tai lieu thanh do")
DANH_MUC = ["Thống kê bài báo đăng tạp chí khoa học trong nước 2021-2025.pdf",
            "Thống kê bài báo đăng tạp chí khoa học quốc tế 2021-2025.pdf"]
DANH_HIEU = {"ts", "ths", "pgs", "gs", "bs", "ds", "cn", "ncs", "dr", "prof", "assoc", "ckii", "cki"}


def _bo_dau(s):
    s = unicodedata.normalize("NFD", s).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def chuan_hoa(ten, quy_tac="A"):
    s = unicodedata.normalize("NFC", str(ten)).lower()
    s = re.sub(r"[.\-,()]", " ", s)
    tu = [t for t in s.split() if _bo_dau(t) not in DANH_HIEU]
    if quy_tac != "A":
        tu = [_bo_dau(t) for t in tu]
    if len(tu) < 2:
        return None
    return tuple(sorted(tu)) if quy_tac == "C" else tuple(tu)


def bai_bao():
    """Trả về danh sách bài báo, mỗi bài là danh sách tên tác giả."""
    ds = []
    for f in DANH_MUC:
        for trang in pymupdf.open(os.path.join(TD, f)):
            for bang in trang.find_tables().tables:
                for hang in bang.extract():
                    if not hang or hang[0] == "Stt":
                        continue
                    stt, tac_gia = (hang[0] or "").strip(), hang[3] or ""
                    dong = [x.strip() for x in tac_gia.split("\n") if x.strip()]
                    if re.fullmatch(r"\d+", stt):
                        ds.append(dong)
                    elif ds and dong:
                        ds[-1].extend(dong)
    return ds


def nhan_su(quy_tac="A"):
    ws = openpyxl.load_workbook(os.path.join(TD, "2026 DS.xlsx"), read_only=True, data_only=True).active
    ten = []
    for hang in list(ws.iter_rows(values_only=True))[1:]:
        ho, t = hang[1], hang[2]
        if ho and t:
            ten.append(chuan_hoa(f"{ho} {t}", quy_tac))
    return ten


def gini(x):
    x = sorted(x)
    n, s = len(x), sum(x)
    return sum((2 * (i + 1) - n - 1) * v for i, v in enumerate(x)) / (n * s)


def tinh(quy_tac="A", bb=None):
    bb = bb or bai_bao()
    ns = nhan_su(quy_tac)
    khoa = set(k for k in ns if k)
    dem = collections.Counter()
    bai_khop = 0
    for tg in bb:
        k_bai = {chuan_hoa(a, quy_tac) for a in tg} & khoa
        if k_bai:
            bai_khop += 1
        dem.update(k_bai)
    co_bai = [dem[k] for k in khoa]
    co = sorted([v for v in co_bai if v > 0], reverse=True)
    top = co[:max(1, round(len(co) * 0.1))]
    return dict(so_bai=len(bb), bai_khop=bai_khop, nhan_su_co_ten=len(ns), ten_chuan_hoa=len(khoa),
                trung_ten=len(ns) - len(khoa), ten_co_bai=len(co), gini=gini(co_bai),
                top10=sum(top) / sum(co))


if __name__ == "__main__":
    ds = bai_bao()
    for qt in "ABC":
        print("Quy tắc", qt)
        for k, v in tinh(qt, ds).items():
            print(f"  {k}: {v:.4f}" if isinstance(v, float) else f"  {k}: {v}")
